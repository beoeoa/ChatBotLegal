"""Build the feature-016 observation baseline from a QA workbook.

The workbook is historical test evidence, not legal ground truth.  This builder
therefore stores only stable questions and aggregate observations.  It never
copies the answer/DOM column into the versioned dataset.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import tempfile
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SHEET_NAME = "Hỏi đáp chi tiết"
REQUIRED_HEADERS = (
    "STT log",
    "Role",
    "Lĩnh vực",
    "Câu hỏi",
    "Câu trả lời trên giao diện",
    "Trạng thái",
    "Thời gian (giây)",
    "Dự phòng",
    "Lỗi",
    "Có nguồn",
    "Mã câu",
    "Ghi chú",
)


def _text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value))).strip()


def _key(value: Any) -> str:
    return _text(value).casefold()


def _flag(value: Any) -> bool:
    return _key(value) in {"1", "true", "yes", "y", "có", "co"}


def _percentile(values: list[float], fraction: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    if len(ordered) == 1:
        return round(ordered[0], 3)
    position = (len(ordered) - 1) * fraction
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    weight = position - lower
    result = ordered[lower] + (ordered[upper] - ordered[lower]) * weight
    return round(result, 3)


def _stable_case_id(value: Any, ordinal: int) -> str:
    raw = _text(value)
    if raw.isdigit():
        return f"web-{int(raw):03d}"
    slug = re.sub(r"[^a-z0-9]+", "-", raw.casefold()).strip("-")
    return f"web-{slug}" if slug else f"web-{ordinal:03d}"


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(handle, "w", encoding="utf-8", newline="\n") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2, sort_keys=False)
            stream.write("\n")
        os.replace(temporary_name, path)
    finally:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)


def build_regression_dataset(
    input_path: str | Path,
    output_path: str | Path,
    *,
    expected_questions: int | None = 100,
    expected_per_domain: int | None = 20,
) -> dict[str, Any]:
    source = Path(input_path).resolve()
    output = Path(output_path).resolve()
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"missing worksheet: {SHEET_NAME}")
        worksheet = workbook[SHEET_NAME]
        rows = worksheet.iter_rows(values_only=True)
        headers = tuple(_text(value) for value in next(rows, ()))
        if headers != REQUIRED_HEADERS:
            raise ValueError(
                "unexpected workbook headers; expected exact QA export contract"
            )
        positions = {header: index for index, header in enumerate(headers)}
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        row_count = 0
        for raw_row in rows:
            if not any(value is not None for value in raw_row):
                continue
            row_count += 1
            row = {
                header: raw_row[index] if index < len(raw_row) else None
                for header, index in positions.items()
            }
            question = _text(row["Câu hỏi"])
            domain = _text(row["Lĩnh vực"])
            if not question or not domain:
                raise ValueError(f"row {row_count + 1} has no question or domain")
            question_key = _key(question)
            grouped[question_key].append(
                {
                    "question": question,
                    "domain": domain,
                    "role": _text(row["Role"]),
                    "status": _key(row["Trạng thái"]),
                    "latency": row["Thời gian (giây)"],
                    "fallback": _flag(row["Dự phòng"]),
                    "error": _flag(row["Lỗi"]),
                    "source_flag": _flag(row["Có nguồn"]),
                    "code": row["Mã câu"],
                    "note": _text(row["Ghi chú"]),
                    "answer_hash": hashlib.sha256(
                        _text(row["Câu trả lời trên giao diện"]).encode("utf-8")
                    ).hexdigest(),
                }
            )
    finally:
        workbook.close()

    if expected_questions is not None and len(grouped) != expected_questions:
        raise ValueError(
            f"expected {expected_questions} unique questions, found {len(grouped)}"
        )

    cases: list[dict[str, Any]] = []
    seen_case_ids: set[str] = set()
    for ordinal, records in enumerate(grouped.values(), start=1):
        domains = {_key(record["domain"]): record["domain"] for record in records}
        if len(domains) != 1:
            raise ValueError("one normalized question is assigned to multiple domains")
        codes = [record["code"] for record in records if _text(record["code"])]
        case_id = _stable_case_id(codes[0] if codes else None, ordinal)
        if case_id in seen_case_ids:
            raise ValueError(f"duplicate case id: {case_id}")
        seen_case_ids.add(case_id)
        latencies: list[float] = []
        for record in records:
            try:
                latencies.append(float(record["latency"]))
            except (TypeError, ValueError):
                pass
        question = records[0]["question"]
        notes = sorted({record["note"] for record in records if record["note"]})
        cases.append(
            {
                "case_id": case_id,
                "question_hash": hashlib.sha256(_key(question).encode("utf-8")).hexdigest(),
                "question": question,
                "domain": records[0]["domain"],
                "observed_roles": sorted(
                    {record["role"] for record in records if record["role"]}
                ),
                "baseline": {
                    "run_count": len(records),
                    "completed_count": sum(
                        record["status"] == "completed" for record in records
                    ),
                    "source_flag_count": sum(
                        record["source_flag"] for record in records
                    ),
                    "fallback_count": sum(record["fallback"] for record in records),
                    "error_count": sum(record["error"] for record in records),
                    "answer_variant_count": len(
                        {record["answer_hash"] for record in records}
                    ),
                    "latency_p50_seconds": _percentile(latencies, 0.5),
                    "latency_p95_seconds": _percentile(latencies, 0.95),
                },
                "workbook_notes": notes,
            }
        )

    cases.sort(key=lambda case: case["case_id"])
    domain_counts = dict(sorted(Counter(case["domain"] for case in cases).items()))
    if expected_per_domain is not None and any(
        count != expected_per_domain for count in domain_counts.values()
    ):
        raise ValueError(
            f"expected {expected_per_domain} questions per domain, found {domain_counts}"
        )

    payload = {
        "schema_version": "1.0",
        "dataset_kind": "historical_observation_not_legal_ground_truth",
        "source": {
            "workbook": source.name,
            "sheet": SHEET_NAME,
            "row_count": row_count,
            "source_sha256": source_hash,
        },
        "unique_question_count": len(cases),
        "domain_counts": domain_counts,
        "cases": cases,
    }
    _write_json_atomic(output, payload)
    return payload


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--expected-questions", type=int, default=100)
    parser.add_argument("--expected-per-domain", type=int, default=20)
    args = parser.parse_args()
    payload = build_regression_dataset(
        args.input,
        args.output,
        expected_questions=args.expected_questions,
        expected_per_domain=args.expected_per_domain,
    )
    print(
        json.dumps(
            {
                "output": str(args.output),
                "questions": payload["unique_question_count"],
                "domains": payload["domain_counts"],
                "source_sha256": payload["source"]["source_sha256"],
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
