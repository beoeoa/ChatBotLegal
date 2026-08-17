from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from io import BytesIO

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter

from api.form_source_resolution import (
    build_pending_candidate,
    classify_effectivity,
    classify_form_effectivity,
    extract_appendix_identifier,
    extract_issuing_instruments,
    extract_strict_form_code,
    extract_strict_form_codes,
    group_occurrences,
    locate_pdf_form_page_range,
    merge_pending_records,
    quarantine_stale_pending_records,
    select_exact_document,
    select_standalone_form_file,
)
from api.official_source_adapters import ordered_fallback_urls
from api.official_source_diagnostics import classify_official_source_failure
from scripts.resolve_three_tier_form_sources import (
    _attach_form_effectivity_evidence,
    _bind_dvc_attachment_rules,
    _bound_group_for_procedure,
    _deterministic_docx_bytes,
    _extract_from_docx_package,
    _extract_from_explicit_pdf_page_range,
    _extract_from_legacy_doc_package,
    _extract_from_pdf_package,
    _filter_prior_resolution_for_active_groups,
    _merge_source_attempts,
    _resumable_prior_resolution,
    _resume_processed_group_ids,
    _validated_content,
    _validated_direct_content,
    _validated_dvc_attachment_content,
    _validated_form_effectivity_rules,
    fetch_official_document,
    resolve_group_artifact,
    resolve_group_artifact_bounded,
)
from scripts.sync_shadow_form_candidates import sync


def test_shadow_candidate_sync_is_idempotent_and_preserves_approved(
    tmp_path,
) -> None:
    queue_path = tmp_path / "queue.json"
    shortlist_path = tmp_path / "shortlist.json"
    queue_path.write_text(
        json.dumps(
            {
                "summary": {},
                "records": [
                    {
                        "id": "three-tier-stale",
                        "approved": False,
                        "runtime_eligible": False,
                    },
                    {
                        "id": "three-tier-approved",
                        "approved": True,
                        "runtime_eligible": True,
                    },
                ],
            }
        ),
        encoding="utf-8",
    )
    shortlist_path.write_text(
        json.dumps(
            {
                "run_id": "shadow-run",
                "records": [
                    {
                        "id": "three-tier-current",
                        "approved": False,
                        "runtime_eligible": False,
                        "sha256": "a" * 64,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    first = sync(shortlist_path, queue_path, tmp_path / "backups")
    second = sync(shortlist_path, queue_path, tmp_path / "backups")
    records = {
        item["id"]: item
        for item in json.loads(queue_path.read_text(encoding="utf-8"))["records"]
    }

    assert first["created"] == 1
    assert second["unchanged"] == 1
    assert records["three-tier-stale"]["catalog_status"] == "quarantined"
    assert records["three-tier-approved"]["approved"] is True
    assert records["three-tier-approved"]["runtime_eligible"] is True
    assert records["three-tier-current"]["approved"] is False
    assert records["three-tier-current"]["runtime_eligible"] is False


def test_resume_filters_stale_groups_after_identity_resolution_changes() -> None:
    prior = {
        "pending_records": [
            {"id": "active-candidate", "three_tier_group_id": "active"},
            {"id": "stale-candidate", "three_tier_group_id": "stale"},
        ],
        "resolved_groups": [
            {"group_id": "active"},
            {"group_id": "stale"},
        ],
        "verified_data_gaps": [
            {"group_id": "active"},
            {"group_id": "stale"},
        ],
        "source_attempts": [
            {"instrument": "01/2024/TT-BTP"},
            {"instrument": "02/2020/TT-BTP"},
        ],
    }

    filtered = _filter_prior_resolution_for_active_groups(
        prior,
        active_group_ids={"active"},
        active_instruments={"01/2024/TT-BTP"},
    )

    assert filtered["pending_records"] == [
        {"id": "active-candidate", "three_tier_group_id": "active"}
    ]
    assert filtered["resolved_groups"] == [{"group_id": "active"}]
    assert filtered["verified_data_gaps"] == [{"group_id": "active"}]
    assert filtered["source_attempts"] == [{"instrument": "01/2024/TT-BTP"}]


def test_resume_keeps_completed_groups_and_deduplicates_source_attempts() -> None:
    payload = {
        "resolved_groups": [{"group_id": "ready-1"}],
        "verified_data_gaps": [
            {"group_id": "gap-1"},
            {"candidate_id": "identity-gap"},
        ],
    }

    assert _resume_processed_group_ids(payload) == {"ready-1", "gap-1"}
    assert _merge_source_attempts(
        [
            {
                "instrument": "02/2024/TT-BYT",
                "status": "found",
                "reason_code": "EXACT_OFFICIAL_DOCUMENT_FOUND",
            }
        ],
        [
            {
                "instrument": "02/2024/TT-BYT",
                "status": "found",
                "reason_code": "EXACT_OFFICIAL_DOCUMENT_FOUND",
            }
        ],
    ) == [
        {
            "instrument": "02/2024/TT-BYT",
            "status": "found",
            "reason_code": "EXACT_OFFICIAL_DOCUMENT_FOUND",
        }
    ]


def test_resume_retries_technical_group_failures_without_reusing_stale_gap() -> None:
    payload = {
        "resolved_groups": [{"group_id": "ready-1"}],
        "verified_data_gaps": [
            {
                "group_id": "retry-timeout",
                "reason_code": "OFFICIAL_ARTIFACT_PROCESSING_TIMEOUT",
            },
            {
                "group_id": "terminal-gap",
                "reason_code": "OFFICIAL_DOCUMENT_IDENTITY_NOT_FOUND",
            },
        ],
    }

    assert _resume_processed_group_ids(payload) == {
        "ready-1",
        "terminal-gap",
    }

    filtered = _filter_prior_resolution_for_active_groups(
        payload,
        active_group_ids={"ready-1", "retry-timeout", "terminal-gap"},
        active_instruments=set(),
    )

    assert filtered["verified_data_gaps"] == [
        {
            "group_id": "terminal-gap",
            "reason_code": "OFFICIAL_DOCUMENT_IDENTITY_NOT_FOUND",
        }
    ]


def test_bounded_artifact_extraction_fails_closed_without_document(tmp_path) -> None:
    artifact, reason = resolve_group_artifact_bounded(
        {"group_id": "group-1", "form_code": "01"},
        {"reason_code": "OFFICIAL_DOCUMENT_IDENTITY_NOT_FOUND"},
        timeout=0.1,
        download_dir=tmp_path,
    )

    assert artifact is None
    assert reason == "OFFICIAL_DOCUMENT_IDENTITY_NOT_FOUND"


def test_group_artifact_extracts_identical_validated_packages_once(
    tmp_path, monkeypatch
) -> None:
    class Client:
        def __init__(self, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    validated_content = b"%PDF-1.7 exact duplicate official package"
    extraction_calls: list[str] = []

    def validated(_client, *, document_id, file_record, cache_dir):
        assert document_id == "official-document-1"
        assert cache_dir.name == "attachments"
        return (
            validated_content,
            f"https://vbpl.vn/{file_record['fileName']}",
            "pdf",
        )

    def extract_pdf(
        content,
        *,
        form_code,
        appendix_identifier,
        download_url,
        download_dir,
        allow_ocr,
    ):
        assert content == validated_content
        assert form_code == "TP-CC-05"
        assert appendix_identifier is None
        assert download_dir == tmp_path
        assert allow_ocr is False
        extraction_calls.append(download_url)
        return {
            "sha256": "a" * 64,
            "file_name": "tp-cc-05.pdf",
        }

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.httpx.Client", Client
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_content", validated
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._extract_from_pdf_package",
        extract_pdf,
    )

    artifact, reason = resolve_group_artifact(
        {"group_id": "group-1", "form_code": "TP-CC-05"},
        {
            "document": {"id": "official-document-1"},
            "files": [
                {"fileName": "appendix-package-a.pdf"},
                {"fileName": "appendix-package-b.pdf"},
            ],
        },
        timeout=1.0,
        download_dir=tmp_path,
    )

    assert artifact == {"sha256": "a" * 64, "file_name": "tp-cc-05.pdf"}
    assert reason == "EXTRACTED_FROM_OFFICIAL_PACKAGE"
    assert extraction_calls == ["https://vbpl.vn/appendix-package-a.pdf"]


def test_group_artifact_keeps_distinct_validated_packages_fail_closed(
    tmp_path, monkeypatch
) -> None:
    class Client:
        def __init__(self, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    extraction_calls: list[bytes] = []

    def validated(_client, *, document_id, file_record, cache_dir):
        assert document_id == "official-document-2"
        assert cache_dir.name == "attachments"
        marker = str(file_record["fileName"]).encode("ascii")
        return b"%PDF-1.7 " + marker, f"https://vbpl.vn/{marker.decode()}", "pdf"

    def extract_pdf(
        content,
        *,
        form_code,
        appendix_identifier,
        download_url,
        download_dir,
        allow_ocr,
    ):
        assert form_code == "TP-CC-08"
        assert appendix_identifier is None
        assert download_dir == tmp_path
        assert allow_ocr is False
        extraction_calls.append(content)
        return {
            "sha256": "a" * 64 if content.endswith(b"a.pdf") else "b" * 64,
            "file_name": download_url.rsplit("/", 1)[-1],
        }

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.httpx.Client", Client
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_content", validated
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._extract_from_pdf_package",
        extract_pdf,
    )

    artifact, reason = resolve_group_artifact(
        {"group_id": "group-2", "form_code": "TP-CC-08"},
        {
            "document": {"id": "official-document-2"},
            "files": [
                {"fileName": "appendix-a.pdf"},
                {"fileName": "appendix-b.pdf"},
            ],
        },
        timeout=1.0,
        download_dir=tmp_path,
    )

    assert artifact is None
    assert reason == "AMBIGUOUS_OFFICIAL_PDF_PACKAGE"
    assert extraction_calls == [
        b"%PDF-1.7 appendix-a.pdf",
        b"%PDF-1.7 appendix-b.pdf",
    ]


def test_group_artifact_prefers_one_explicit_direct_official_package(
    tmp_path, monkeypatch
) -> None:
    class Client:
        def __init__(self, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    direct_url = "https://datafiles.chinhphu.vn/cpp/files/direct.pdf"

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.httpx.Client", Client
    )
    attachment_calls = []

    def validated_attachment(*_args, **_kwargs):
        attachment_calls.append("attachment")
        return (
            b"%PDF-1.7 attachment",
            "https://vbpl.vn/attachment.pdf",
            "pdf",
        )

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_content",
        validated_attachment,
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_direct_content",
        lambda *_args, **_kwargs: (b"%PDF-1.7 direct", direct_url, "pdf"),
    )

    def extract_pdf(content, **_kwargs):
        direct = content.endswith(b"direct")
        return {
            "sha256": ("d" if direct else "a") * 64,
            "file_name": "direct.pdf" if direct else "attachment.pdf",
        }

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._extract_from_pdf_package",
        extract_pdf,
    )

    artifact, reason = resolve_group_artifact(
        {
            "group_id": "group-explicit-direct",
            "form_code": "21",
            "official_download_urls": [direct_url],
        },
        {
            "document": {"id": "official-document-direct"},
            "files": [{"fileName": "attachment.pdf"}],
        },
        timeout=1.0,
        download_dir=tmp_path,
    )

    assert artifact is not None
    assert artifact["sha256"] == "d" * 64
    assert artifact["file_name"] == "direct.pdf"
    assert artifact["source_page_url"] == direct_url
    assert artifact["source_retrieval_url"] == direct_url
    assert artifact["source_retrieval_method"] == "GET"
    assert artifact["provenance_kind"] == "official_direct_attachment"
    assert reason == "EXTRACTED_FROM_DIRECT_OFFICIAL_PACKAGE"
    assert attachment_calls == []


def test_group_artifact_extracts_checksum_pinned_curated_page_range(
    tmp_path, monkeypatch
) -> None:
    class Client:
        def __init__(self, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    writer = PdfWriter()
    for _index in range(3):
        writer.add_blank_page(width=72, height=72)
    package = BytesIO()
    writer.write(package)
    content = package.getvalue()
    landing_url = "https://cathai.haiphong.gov.vn/qd52"
    download_url = "https://cdn.haiphong.gov.vn/qd52.pdf"

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.httpx.Client", Client
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_direct_content",
        lambda *_args, **_kwargs: (content, download_url, "pdf"),
    )

    artifact, reason = resolve_group_artifact(
        {
            "group_id": "qd52-form-01",
            "form_code": "01",
            "appendix_identifier": "04",
            "procedure_ids": ["1.115594", "1.115595"],
        },
        {
            "document": {
                "id": "hai-phong-qd52-2026",
                "detailUrl": landing_url,
                "publisher": "Ủy ban nhân dân thành phố Hải Phòng",
            },
            "files": [],
            "explicit_form_page_bindings": [
                {
                    "form_code": "01",
                    "appendix_identifier": "04",
                    "procedure_ids": ["1.115594", "1.115595"],
                    "source_page_url": landing_url,
                    "official_download_url": download_url,
                    "source_package_sha256": hashlib.sha256(content).hexdigest(),
                    "source_package_size_bytes": len(content),
                    "source_pages_zero_based": [1, 2],
                }
            ],
        },
        timeout=1.0,
        download_dir=tmp_path,
    )

    assert reason == "EXTRACTED_FROM_CURATED_OFFICIAL_PACKAGE_PAGE_RANGE"
    assert artifact is not None
    assert artifact["source_page_url"] == landing_url
    assert artifact["source_retrieval_url"] == download_url
    assert artifact["source_pages_zero_based"] == [1, 2]
    assert artifact["source_package_sha256"] == hashlib.sha256(content).hexdigest()
    assert artifact["extraction"]["kind"] == "curated_official_package_page_range"
    from pypdf import PdfReader

    assert len(PdfReader(tmp_path / artifact["file_name"]).pages) == 2


def test_group_artifact_rejects_curated_package_checksum_drift(
    tmp_path, monkeypatch
) -> None:
    class Client:
        def __init__(self, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    package = BytesIO()
    writer.write(package)
    content = package.getvalue()
    landing_url = "https://cathai.haiphong.gov.vn/qd52"
    download_url = "https://cdn.haiphong.gov.vn/qd52.pdf"

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.httpx.Client", Client
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_direct_content",
        lambda *_args, **_kwargs: (content, download_url, "pdf"),
    )

    artifact, reason = resolve_group_artifact(
        {
            "group_id": "qd52-form-01",
            "form_code": "01",
            "appendix_identifier": "04",
            "procedure_ids": ["1.115594"],
        },
        {
            "document": {"id": "hai-phong-qd52-2026"},
            "files": [],
            "explicit_form_page_bindings": [
                {
                    "form_code": "01",
                    "appendix_identifier": "04",
                    "procedure_ids": ["1.115594"],
                    "source_page_url": landing_url,
                    "official_download_url": download_url,
                    "source_package_sha256": "0" * 64,
                    "source_package_size_bytes": len(content),
                    "source_pages_zero_based": [0],
                }
            ],
        },
        timeout=1.0,
        download_dir=tmp_path,
    )

    assert artifact is None
    assert reason == "CURATED_OFFICIAL_PACKAGE_CHECKSUM_DRIFT"


def test_explicit_pdf_page_range_rejects_noncontiguous_pages(tmp_path) -> None:
    writer = PdfWriter()
    for _index in range(3):
        writer.add_blank_page(width=72, height=72)
    package = BytesIO()
    writer.write(package)

    assert _extract_from_explicit_pdf_page_range(
        package.getvalue(),
        form_code="01",
        appendix_identifier="4",
        source_pages_zero_based=[0, 2],
        download_url="https://cdn.haiphong.gov.vn/qd52.pdf",
        download_dir=tmp_path,
    ) is None


def test_group_artifact_marks_duplicate_explicit_direct_package_as_preferred(
    tmp_path, monkeypatch
) -> None:
    class Client:
        def __init__(self, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    direct_url = "https://datafiles.chinhphu.vn/cpp/files/direct.pdf"
    direct_content = b"%PDF-1.7 direct"
    other_content = b"%PDF-1.7 other"

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.httpx.Client", Client
    )

    def validated(_client, *, file_record, **_kwargs):
        if file_record["fileName"] == "direct.pdf":
            return direct_content, direct_url, "pdf"
        return other_content, "https://vbpl.vn/other.pdf", "pdf"

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_content",
        validated,
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_direct_content",
        lambda *_args, **_kwargs: (direct_content, direct_url, "pdf"),
    )

    def extract_pdf(content, **_kwargs):
        direct = content == direct_content
        return {
            "sha256": ("d" if direct else "a") * 64,
            "file_name": "direct.pdf" if direct else "other.pdf",
            "source_package_sha256": hashlib.sha256(content).hexdigest(),
        }

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._extract_from_pdf_package",
        extract_pdf,
    )

    artifact, reason = resolve_group_artifact(
        {
            "group_id": "group-duplicate-explicit-direct",
            "form_code": "21",
            "official_download_urls": [direct_url],
        },
        {
            "document": {"id": "official-document-direct"},
            "files": [
                {"fileName": "direct.pdf"},
                {"fileName": "other.pdf"},
            ],
        },
        timeout=1.0,
        download_dir=tmp_path,
    )

    assert artifact is not None
    assert artifact["sha256"] == "d" * 64
    assert reason == "EXTRACTED_FROM_DIRECT_OFFICIAL_PACKAGE"


def test_extract_strict_form_code_accepts_real_codes() -> None:
    cases = {
        "Mẫu CT01 ban hành kèm theo Thông tư số 53/2025/TT-BCA": "CT01",
        "mẫu NA17 dùng trong lĩnh vực xuất nhập cảnh": "NA17",
        "Mẫu M01a - Đơn đề nghị": "M01A",
        "Mẫu số 01: Tờ khai": "01",
        "Mẫu số 09/ĐK": "09/ĐK",
    }

    for text, expected in cases.items():
        assert extract_strict_form_code(text) == expected


def test_extract_all_form_codes_detects_ambiguous_same_page() -> None:
    assert extract_strict_form_codes("Mẫu số 01\n...\nMẫu số 02") == ["01", "02"]


def test_extract_all_form_codes_expands_compact_lists_and_compound_codes() -> None:
    assert extract_strict_form_codes("Mẫu số 06, 07 và 08 ban hành kèm theo") == [
        "06",
        "07",
        "08",
    ]
    assert extract_strict_form_codes("Mẫu số 18 và Mẫu số 19") == ["18", "19"]
    assert extract_strict_form_codes(
        "Mẫu TP-TSCC-01a và TP-TSCC-01b"
    ) == ["TP-TSCC-01A", "TP-TSCC-01B"]
    assert extract_strict_form_codes("theo mẫu AP02; Mẫu DC02") == ["AP02", "DC02"]
    assert extract_strict_form_codes("Mẫu TP-CC-01ban hành kèm theo") == [
        "TP-CC-01"
    ]
    assert extract_strict_form_codes("Mấu số 16 (QĐ)") == ["16"]


def test_extract_issuing_instrument_tolerates_portal_spacing() -> None:
    assert extract_issuing_instruments(
        "ban hành kèm theo Thông tư số 40 /2013/TT-BCT"
    ) == ["40/2013/TT-BCT"]
    assert extract_issuing_instruments("Nghị định số 131/2021/NĐ- CP") == [
        "131/2021/NĐ-CP"
    ]


def test_extract_strict_form_code_rejects_natural_language_false_positives() -> None:
    for text in (
        "mẫu hộ gia đình",
        "mẫu quy định tại phụ lục",
        "mẫu tại cơ quan tiếp nhận",
        "biểu mẫu điện tử tương tác",
    ):
        assert extract_strict_form_code(text) is None


def test_extract_issuing_instrument_uses_explicit_issued_with_phrase() -> None:
    text = "Mẫu CT01 ban hành kèm theo Thông tư số 53/2025/TT-BCA"

    assert extract_issuing_instruments(text) == ["53/2025/TT-BCA"]
    assert extract_issuing_instruments(
        "ban hành kèm theo Nghị quyết số 66.18/2026/NQ-CP"
    ) == ["66.18/2026/NQ-CP"]


def test_extract_issuing_instrument_reports_multiple_explicit_documents() -> None:
    text = (
        "Mẫu số 01 ban hành kèm theo Thông tư 01/2024/TT-BTP; "
        "được thay thế theo Thông tư 02/2025/TT-BTP"
    )

    assert extract_issuing_instruments(text) == [
        "01/2024/TT-BTP",
        "02/2025/TT-BTP",
    ]


def test_group_occurrences_prefers_explicit_issued_with_instrument_over_amendments() -> None:
    result = group_occurrences(
        [
            {
                "candidate_id": "occ-na17",
                "procedure_id": "1.000253",
                "form_name": (
                    "Phiếu khai báo tạm trú (mẫu NA17) ban hành kèm theo "
                    "Thông tư số 04/2015/TT-BCA, được sửa đổi bởi "
                    "Thông tư số 57/2020/TT-BCA và Thông tư số 22/2023/TT-BCA"
                ),
                "issuing_instruments": [
                    "04/2015/TT-BCA",
                    "57/2020/TT-BCA",
                    "22/2023/TT-BCA",
                ],
            }
        ]
    )

    assert result["unresolved"] == []
    assert result["groups"][0]["form_code"] == "NA17"
    assert result["groups"][0]["issuing_instrument"] == "04/2015/TT-BCA"


def test_group_occurrences_deduplicates_by_code_and_issuing_instrument() -> None:
    occurrences = [
        {
            "candidate_id": "occ-a",
            "procedure_id": "1.001",
            "form_name": "Mẫu CT01 ban hành kèm theo Thông tư 53/2025/TT-BCA",
            "executing_level": "province",
        },
        {
            "candidate_id": "occ-b",
            "procedure_id": "1.002",
            "form_name": "Mẫu CT01 - Thông tư số 53/2025/TT-BCA",
            "issuing_instruments": ["53/2025/TT-BCA"],
            "executing_level": "province",
        },
        {
            "candidate_id": "occ-c",
            "procedure_id": "1.003",
            "form_name": "Mẫu CT01 ban hành kèm theo Thông tư 55/2021/TT-BCA",
            "executing_level": "commune",
        },
    ]

    result = group_occurrences(occurrences)

    assert len(result["groups"]) == 2
    current = next(
        item
        for item in result["groups"]
        if item["issuing_instrument"] == "53/2025/TT-BCA"
    )
    assert current["form_code"] == "CT01"
    assert current["procedure_ids"] == ["1.001", "1.002"]
    assert current["occurrence_ids"] == ["occ-a", "occ-b"]
    assert current["executing_level"] == "province"
    assert result["unresolved"] == []


def test_bound_group_keeps_aggregated_executing_level_when_binding_is_blank() -> None:
    bound = _bound_group_for_procedure(
        {
            "domain": "an_sinh_y_te_giao_duc",
            "source_tier": "central",
            "executing_level": "province",
            "procedure_metadata": {
                "1.013868": {
                    "domain": "an_sinh_y_te_giao_duc",
                    "source_tier": "central",
                }
            },
        },
        "1.013868",
    )

    assert bound["executing_level"] == "province"


def test_completed_probe_is_not_reused_as_a_resume_checkpoint() -> None:
    completed = {
        "run_id": "completed-run",
        "group_progress_complete": True,
        "resolved_groups": [{"group_id": "group-1"}],
        "pending_records": [{"id": "candidate-1"}],
    }
    interrupted = {
        **completed,
        "group_progress_complete": False,
    }

    assert _resumable_prior_resolution(completed) == {}
    assert _resumable_prior_resolution(interrupted)["run_id"] == "completed-run"


def test_group_occurrences_prefers_explicit_form_code_when_title_lists_alternatives() -> None:
    title = (
        "Giấy đề nghị theo Mẫu TP-TSCC-01A áp dụng trường hợp tự liên hệ "
        "hoặc Mẫu TP-TSCC-01B áp dụng trường hợp bị từ chối nhận tập sự, "
        "ban hành kèm theo Thông tư 06/2025/TT-BTP"
    )

    result = group_occurrences(
        [
            {
                "candidate_id": "occurrence-a",
                "procedure_id": "1.013807",
                "form_name": title,
                "form_code": "TP-TSCC-01A",
                "issuing_instrument": "06/2025/TT-BTP",
            },
            {
                "candidate_id": "occurrence-b",
                "procedure_id": "1.013807",
                "form_name": title,
                "form_code": "TP-TSCC-01B",
                "issuing_instrument": "06/2025/TT-BTP",
            },
        ]
    )

    assert {item["form_code"] for item in result["groups"]} == {
        "TP-TSCC-01A",
        "TP-TSCC-01B",
    }
    assert result["unresolved"] == []


def test_group_occurrences_splits_repeated_code_by_exact_appendix() -> None:
    result = group_occurrences(
        [
            {
                "candidate_id": "occ-appendix-i",
                "procedure_id": "1.001",
                "form_name": (
                    "Mẫu 01 Phụ lục I ban hành kèm theo "
                    "Nghị định 96/2023/NĐ-CP"
                ),
            },
            {
                "candidate_id": "occ-appendix-ii",
                "procedure_id": "1.002",
                "form_name": (
                    "Mẫu 01 Phụ lục II ban hành kèm theo "
                    "Nghị định 96/2023/NĐ-CP"
                ),
            },
        ]
    )

    assert result["unresolved"] == []
    assert len(result["groups"]) == 2
    assert {item["appendix_identifier"] for item in result["groups"]} == {
        "I",
        "II",
    }
    assert len({item["group_id"] for item in result["groups"]}) == 2


def test_group_occurrences_extracts_appendix_number_after_so() -> None:
    result = group_occurrences(
        [
            {
                "candidate_id": "qd52-01",
                "procedure_id": "1.115594",
                "form_name": "Mẫu 01 Phụ lục số 04",
                "issuing_instruments": ["52/2026/QĐ-UBND"],
            }
        ]
    )

    assert extract_appendix_identifier("Phụ lục số 04") == "04"
    assert result["unresolved"] == []
    assert result["groups"][0]["appendix_identifier"] == "04"


def test_group_occurrences_never_merges_without_code_or_instrument() -> None:
    result = group_occurrences(
        [
            {
                "candidate_id": "occ-a",
                "procedure_id": "1.001",
                "form_name": "Tờ khai theo quy định",
            }
        ]
    )

    assert result["groups"] == []
    assert result["unresolved"][0]["reason_code"] == "FORM_IDENTITY_UNRESOLVED"


def test_select_exact_document_rejects_fuzzy_number() -> None:
    records = [
        {"id": 1, "docNum": "153/2025/TT-BCA"},
        {"id": 2, "docNum": "53/2025/TT-BCA"},
        {"id": 3, "docNum": "53/2025/TT-BTP"},
    ]

    assert select_exact_document(records, "53/2025/TT-BCA")["id"] == 2
    assert select_exact_document(records, "53/2025/TT-BXD") is None


def test_select_exact_document_accepts_official_number_label() -> None:
    records = [{"id": "doc-154", "docNum": "Số: 154/2024/NĐ-CP"}]

    assert select_exact_document(records, "154/2024/NĐ-CP")["id"] == "doc-154"


def test_effectivity_is_fail_closed() -> None:
    assert classify_effectivity(
        {"effStatus": "Còn hiệu lực", "effFrom": "2025-07-01", "effTo": None},
        legal_as_of="2026-07-27",
    )["eligible"] is True


def test_partial_document_needs_exact_appendix_effectivity_evidence() -> None:
    document = {
        "effStatus": "Hết hiệu lực một phần",
        "effFrom": "2025-07-01",
        "appendix_effectivity": [
            {
                "appendix_identifier": "II",
                "status": "active",
                "official_source_url": "https://vbpl.vn/van-ban/example",
                "verified_as_of": "2026-07-27",
            }
        ],
    }

    assert classify_form_effectivity(
        document,
        legal_as_of="2026-07-27",
        appendix_identifier="I",
    )["reason_code"] == "PARTIAL_EFFECTIVITY_REQUIRES_REVIEW"
    exact = classify_form_effectivity(
        document,
        legal_as_of="2026-07-27",
        appendix_identifier="II",
    )
    assert exact["eligible"] is True
    assert exact["reason_code"] == "APPENDIX_EFFECTIVITY_VERIFIED"


def test_partial_document_accepts_exact_form_code_effectivity_evidence() -> None:
    decision = classify_form_effectivity(
        {
            "docNum": "05/2025/TT-BTP",
            "effStatus": "Hết hiệu lực một phần",
            "effFrom": "2025-07-01",
            "appendix_effectivity": [
                {
                    "target_form_code": "TP-CC-06",
                    "status": "active",
                    "official_source_url": (
                        "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=178734"
                    ),
                    "verified_as_of": "2026-07-29",
                }
            ],
        },
        legal_as_of="2026-07-29",
        appendix_identifier=None,
        form_code="TP-CC-06",
    )

    assert decision["eligible"] is True
    assert decision["reason_code"] == "FORM_EFFECTIVITY_VERIFIED"


def test_effectivity_sidecar_requires_exact_safe_replacement_scope() -> None:
    payload = {
        "form_effectivity": [
            {
                "evidence_id": "tp-cc-06-active",
                "issuing_instrument": "05/2025/TT-BTP",
                "target_form_code": "TP-CC-06",
                "status": "active",
                "verified_as_of": "2026-07-29",
                "official_source_url": (
                    "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=178734"
                ),
                "history_source_url": (
                    "https://vbpl.vn/botuphap/Pages/vbpq-lichsu.aspx?ItemID=178491"
                ),
                "evidence_basis": "explicit_replacement_scope",
                "relation_edges": [
                    {
                        "relation": "replaces_other_forms_only",
                        "modifying_document_number": "11/2025/TT-BTP",
                        "effective_from": "2025-07-01",
                        "affected_form_codes": [
                            "TP-CC-01",
                            "TP-CC-02",
                            "TP-CC-03",
                            "TP-CC-04",
                        ],
                        "official_source_url": (
                            "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=178734"
                        ),
                    }
                ],
            },
            {
                "evidence_id": "unsafe-target-replaced",
                "issuing_instrument": "05/2025/TT-BTP",
                "target_form_code": "TP-CC-06",
                "status": "active",
                "verified_as_of": "2026-07-29",
                "official_source_url": "https://example.com/not-official",
                "evidence_basis": "explicit_replacement_scope",
                "relation_edges": [
                    {
                        "relation": "replaces_other_forms_only",
                        "modifying_document_number": "11/2025/TT-BTP",
                        "affected_form_codes": ["TP-CC-06"],
                        "official_source_url": "https://example.com/not-official",
                    }
                ],
            },
        ]
    }

    rules, rejections = _validated_form_effectivity_rules(
        payload,
        legal_as_of="2026-07-29",
    )
    document = _attach_form_effectivity_evidence(
        {
            "docNum": "05/2025/TT-BTP",
            "effStatus": "Hết hiệu lực một phần",
            "effFrom": "2025-07-01",
        },
        rules,
    )

    assert [item["evidence_id"] for item in rules] == ["tp-cc-06-active"]
    assert rejections == [
        {
            "evidence_id": "unsafe-target-replaced",
            "reason_code": "FORM_EFFECTIVITY_SOURCE_NOT_OFFICIAL",
        }
    ]
    assert document["appendix_effectivity"] == [rules[0]]


def test_effectivity_sidecar_accepts_only_exact_current_official_form_attachment() -> None:
    form_scope_page = (
        "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=179354"
    )
    attachment_page = (
        "https://vbpl.vn/TW/Pages/vbpq-van-ban-goc.aspx?ItemID=179354"
    )
    current_status_page = (
        "https://vbpl.vn/TW/Pages/vbpq-thuoctinh.aspx?ItemID=179354"
    )
    ct01_download = (
        "https://vbpl-bientap-gateway.moj.gov.vn/api/qtdc/public/doc/minio/"
        "buckets/vbpl/179354/1.%20M%E1%BA%ABu%20CT01%20ban%20h%C3%A0nh%20"
        "k%C3%A8m%20theo%20Th%C3%B4ng%20t%C6%B0%2053.doc/download"
    )
    payload = {
        "form_effectivity": [
            {
                "evidence_id": "tt-53-2025-bca-ct01-current",
                "issuing_instrument": "53/2025/TT-BCA",
                "target_form_code": "CT01",
                "status": "current",
                "verified_as_of": "2026-07-29",
                "official_source_url": form_scope_page,
                "official_attachment_source_url": attachment_page,
                "current_status_source_url": current_status_page,
                "official_form_file_name": (
                    "1. Mẫu CT01 ban hành kèm theo Thông tư 53.doc"
                ),
                "official_form_download_url": ct01_download,
                "evidence_basis": "official_current_form_attachment",
            },
            {
                "evidence_id": "tt-53-2025-bca-wrong-code",
                "issuing_instrument": "53/2025/TT-BCA",
                "target_form_code": "CT01",
                "status": "current",
                "verified_as_of": "2026-07-29",
                "official_source_url": form_scope_page,
                "official_attachment_source_url": attachment_page,
                "current_status_source_url": current_status_page,
                "official_form_file_name": (
                    "2. Mẫu CT02 ban hành kèm theo Thông tư 53.doc"
                ),
                "official_form_download_url": ct01_download,
                "evidence_basis": "official_current_form_attachment",
            },
            {
                "evidence_id": "tt-53-2025-bca-unofficial-status",
                "issuing_instrument": "53/2025/TT-BCA",
                "target_form_code": "CT01",
                "status": "current",
                "verified_as_of": "2026-07-29",
                "official_source_url": form_scope_page,
                "official_attachment_source_url": attachment_page,
                "current_status_source_url": "https://example.com/current",
                "official_form_file_name": (
                    "1. Mẫu CT01 ban hành kèm theo Thông tư 53.doc"
                ),
                "official_form_download_url": ct01_download,
                "evidence_basis": "official_current_form_attachment",
            },
            {
                "evidence_id": "tt-53-2025-bca-missing-download",
                "issuing_instrument": "53/2025/TT-BCA",
                "target_form_code": "CT01",
                "status": "current",
                "verified_as_of": "2026-07-29",
                "official_source_url": form_scope_page,
                "official_attachment_source_url": attachment_page,
                "current_status_source_url": current_status_page,
                "official_form_file_name": (
                    "1. Mẫu CT01 ban hành kèm theo Thông tư 53.doc"
                ),
                "official_form_download_url": "",
                "evidence_basis": "official_current_form_attachment",
            },
        ]
    }

    rules, rejections = _validated_form_effectivity_rules(
        payload,
        legal_as_of="2026-07-29",
    )

    assert [item["evidence_id"] for item in rules] == [
        "tt-53-2025-bca-ct01-current"
    ]
    assert rejections == [
        {
            "evidence_id": "tt-53-2025-bca-wrong-code",
            "reason_code": "FORM_EFFECTIVITY_ATTACHMENT_CODE_MISMATCH",
        },
        {
            "evidence_id": "tt-53-2025-bca-unofficial-status",
            "reason_code": "FORM_EFFECTIVITY_CURRENT_STATUS_SOURCE_NOT_OFFICIAL",
        },
        {
            "evidence_id": "tt-53-2025-bca-missing-download",
            "reason_code": "FORM_EFFECTIVITY_ATTACHMENT_INCOMPLETE",
        },
    ]


def _current_dvc_attachment_rule() -> dict:
    procedure_url = (
        "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/"
        "019d2bff-2d36-7533-81f7-2f3380a96f1b"
    )
    return {
        "evidence_id": "nd-91-2016-form-05-current-dvc",
        "issuing_instrument": "91/2016/NĐ-CP",
        "target_form_code": "05",
        "appendix_identifier": "I",
        "status": "current",
        "effective_from": "2026-07-01",
        "effective_to": "2027-02-28",
        "verified_as_of": "2026-07-29",
        "official_source_url": procedure_url,
        "history_source_url": (
            "https://vbpl.vn/tw/Pages/vbpq-lichsu.aspx?ItemID=112021"
        ),
        "current_status_source_url": (
            "https://dichvucong.gov.vn/quyet-dinh-cong-bo/"
            "019eb146-cf41-72c8-9f55-1d3df27dd74a"
        ),
        "publication_decision_number": "1684/QĐ-BYT",
        "procedure_bindings": [
            {
                "procedure_id": "1.013868",
                "official_source_url": procedure_url,
                "binding_text": (
                    "Văn bản đề nghị đăng ký lưu hành bổ sung theo Mẫu số 05 "
                    "tại Phụ lục I ban hành kèm theo Nghị định số "
                    "91/2016/NĐ-CP, Nghị định số 155/2018/NĐ-CP và "
                    "Nghị định số 129/2024/NĐ-CP."
                ),
            }
        ],
        "canonical_artifact_attachment": {
            "api_url": (
                "https://dichvucong.gov.vn/api/v1/submitting/"
                "preview-attachment"
            ),
            "attachment_id": "019e2aed-9d64-73ee-a43e-245d6a2a06e6",
            "file_name": "maudon.docx",
            "sha256": (
                "f5974d2ded3229858901d9680f7c0fed55e7e8f68bd188ebd49f99d363a54c55"
            ),
            "size_bytes": 24924,
            "referer_url": procedure_url,
            "source_page_url": procedure_url,
        },
        "evidence_basis": "official_current_dvc_attachment",
    }


def test_effectivity_sidecar_accepts_generic_dvc_filename_only_with_exact_binding() -> None:
    valid = _current_dvc_attachment_rule()
    wrong_endpoint = json.loads(json.dumps(valid))
    wrong_endpoint["evidence_id"] = "nd-91-2016-form-05-wrong-endpoint"
    wrong_endpoint["canonical_artifact_attachment"]["api_url"] = (
        "https://dichvucong.gov.vn/api/v1/submitting/other"
    )
    wrong_binding = json.loads(json.dumps(valid))
    wrong_binding["evidence_id"] = "nd-91-2016-form-05-wrong-binding"
    wrong_binding["procedure_bindings"][0]["binding_text"] = (
        "Mẫu số 06 tại Phụ lục I ban hành kèm theo "
        "Nghị định số 91/2016/NĐ-CP."
    )
    missing_integrity = json.loads(json.dumps(valid))
    missing_integrity["evidence_id"] = "nd-91-2016-form-05-missing-integrity"
    missing_integrity["canonical_artifact_attachment"]["sha256"] = ""

    rules, rejections = _validated_form_effectivity_rules(
        {
            "form_effectivity": [
                valid,
                wrong_endpoint,
                wrong_binding,
                missing_integrity,
            ]
        },
        legal_as_of="2026-07-29",
    )

    assert [item["evidence_id"] for item in rules] == [
        "nd-91-2016-form-05-current-dvc"
    ]
    assert rejections == [
        {
            "evidence_id": "nd-91-2016-form-05-wrong-endpoint",
            "reason_code": "FORM_EFFECTIVITY_DVC_ATTACHMENT_ENDPOINT_INVALID",
        },
        {
            "evidence_id": "nd-91-2016-form-05-wrong-binding",
            "reason_code": "FORM_EFFECTIVITY_DVC_BINDING_IDENTITY_MISMATCH",
        },
        {
            "evidence_id": "nd-91-2016-form-05-missing-integrity",
            "reason_code": "FORM_EFFECTIVITY_DVC_ATTACHMENT_INTEGRITY_INCOMPLETE",
        },
    ]


def test_dvc_attachment_rule_requires_all_group_procedure_bindings() -> None:
    rule = _current_dvc_attachment_rule()
    groups = [
        {
            "group_id": "group-05",
            "form_code": "05",
            "issuing_instrument": "91/2016/NĐ-CP",
            "appendix_identifier": "I",
            "procedure_ids": ["1.013868", "1.013895"],
        }
    ]

    issues = _bind_dvc_attachment_rules(groups, [rule])

    assert issues == [
        {
            "group_id": "group-05",
            "reason_code": "OFFICIAL_DVC_PROCEDURE_BINDING_MISMATCH",
        }
    ]
    assert groups[0]["official_dvc_attachment_required"] is True
    assert (
        groups[0]["official_dvc_attachment_reason_code"]
        == "OFFICIAL_DVC_PROCEDURE_BINDING_MISMATCH"
    )
    assert groups[0]["official_dvc_attachment_requests"] == []


def test_effectivity_sidecar_accepts_only_disjoint_expired_provision_scope() -> None:
    modifying_text = (
        "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=180280"
    )
    parent_text = (
        "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=178434"
    )
    parent_history = (
        "https://vbpl.vn/bonoivu/Pages/vbpq-lichsu.aspx?ItemID=178434"
    )
    package_download = (
        "https://vbpl-bientap-gateway.moj.gov.vn/api/qtdc/public/doc/minio/"
        "buckets/vbpl/178434/09.2025.TT-NV.docx/download"
    )
    base = {
        "issuing_instrument": "09/2025/TT-BNV",
        "target_form_code": "03",
        "status": "active",
        "verified_as_of": "2026-07-29",
        "official_source_url": modifying_text,
        "history_source_url": parent_history,
        "form_scope_source_url": parent_text,
        "target_provisions": ["11"],
        "official_form_package_file_name": "09.2025.TT-NV.docx",
        "official_form_package_download_url": package_download,
        "official_form_package_sha256": (
            "1eabca2e12d27dce2fe7391201baf916c268c4c69890cdd32af1245ed402e380"
        ),
        "official_form_package_size": 128426,
        "evidence_basis": "explicit_partial_provision_scope",
        "relation_edges": [
            {
                "relation": "expires_other_provisions_only",
                "modifying_document_number": "15/2025/TT-BNV",
                "effective_from": "2025-08-05",
                "affected_provisions": ["3", "4", "5"],
                "official_source_url": modifying_text,
            },
            {
                "relation": "repeals_other_provisions_only",
                "modifying_document_number": "16/2025/TT-BNV",
                "effective_from": "2025-08-06",
                "affected_provisions": ["6"],
                "official_source_url": modifying_text,
            }
        ],
    }
    payload = {
        "form_effectivity": [
            {"evidence_id": "tt-09-form-03-active", **base},
            {
                "evidence_id": "tt-09-form-03-overlap",
                **base,
                "target_provisions": ["5", "11"],
            },
            {
                "evidence_id": "tt-09-form-03-unofficial-scope",
                **base,
                "form_scope_source_url": "https://example.com/form-scope",
            },
            {
                "evidence_id": "tt-09-form-03-missing-history",
                **base,
                "history_source_url": "",
            },
            {
                "evidence_id": "tt-09-form-03-unofficial-package",
                **base,
                "official_form_package_download_url": (
                    "https://example.com/09.2025.TT-NV.docx"
                ),
            },
        ]
    }

    rules, rejections = _validated_form_effectivity_rules(
        payload,
        legal_as_of="2026-07-29",
    )

    assert [item["evidence_id"] for item in rules] == [
        "tt-09-form-03-active"
    ]
    assert rejections == [
        {
            "evidence_id": "tt-09-form-03-overlap",
            "reason_code": "FORM_EFFECTIVITY_PROVISION_SCOPE_UNSAFE",
        },
        {
            "evidence_id": "tt-09-form-03-unofficial-scope",
            "reason_code": "FORM_EFFECTIVITY_PROVISION_SOURCE_NOT_OFFICIAL",
        },
        {
            "evidence_id": "tt-09-form-03-missing-history",
            "reason_code": "FORM_EFFECTIVITY_PROVISION_HISTORY_MISSING",
        },
        {
            "evidence_id": "tt-09-form-03-unofficial-package",
            "reason_code": "FORM_EFFECTIVITY_PACKAGE_SOURCE_NOT_OFFICIAL",
        },
    ]


def test_effectivity_sidecar_accepts_only_exact_suspended_repeal_scope() -> None:
    parent_text = (
        "https://vbpl.vn/van-ban/chi-tiet/nghi-dinh-so-148-2025-nd-cp-"
        "quy-dinh-ve-phan-quyen-phan-cap-trong-linh-vuc-y-te--178254"
    )
    parent_history = (
        "https://vbpl.vn/tw/Pages/vbpq-lichsu.aspx?ItemID=178254"
    )
    package_download = (
        "https://vbpl-bientap-gateway.moj.gov.vn/api/qtdc/public/doc/minio/"
        "buckets/vbpl/178254/VanBanGoc_148-2025-nd-cp-trang-5.pdf/download"
    )
    repealer_source = (
        "https://congbao.chinhphu.vn/van-ban/"
        "nghi-dinh-so-46-2026-nd-cp-468886.htm"
    )
    suspension_source = (
        "https://vanban.chinhphu.vn/?classid=1&docid=217667&pageid=27160"
    )
    base = {
        "issuing_instrument": "148/2025/NĐ-CP",
        "target_form_code": "03",
        "status": "active",
        "verified_as_of": "2026-07-29",
        "official_source_url": suspension_source,
        "history_source_url": parent_history,
        "form_scope_source_url": parent_text,
        "target_provisions": ["28.1"],
        "official_form_package_file_name": (
            "VanBanGoc_148-2025-nd-cp-trang-5.pdf"
        ),
        "official_form_package_download_url": package_download,
        "official_form_package_sha256": (
            "6ef20f025c1dd6fd37ced4d38e4cdbf01354cf19aa1d666e0860fa57fe8d73bc"
        ),
        "official_form_package_size": 8955768,
        "evidence_basis": "explicit_suspended_repeal_scope",
        "relation_edges": [
            {
                "relation": "repeals_target_provisions",
                "modifying_document_number": "46/2026/NĐ-CP",
                "effective_from": "2026-01-26",
                "affected_provisions": ["25", "28.1", "30"],
                "official_source_url": repealer_source,
            },
            {
                "relation": "suspends_repealing_instrument",
                "modifying_document_number": "15/2026/NQ-CP",
                "suspended_document_number": "46/2026/NĐ-CP",
                "effective_from": "2026-04-06",
                "suspension_status": "active",
                "official_source_url": suspension_source,
            },
        ],
    }
    payload = {
        "form_effectivity": [
            {"evidence_id": "nd-148-form-03-active", **base},
            {
                "evidence_id": "nd-148-form-03-no-target-overlap",
                **base,
                "target_provisions": ["20"],
            },
            {
                "evidence_id": "nd-148-form-03-wrong-suspended-document",
                **base,
                "relation_edges": [
                    base["relation_edges"][0],
                    {
                        **base["relation_edges"][1],
                        "suspended_document_number": "15/2018/NĐ-CP",
                    },
                ],
            },
            {
                "evidence_id": "nd-148-form-03-inactive-suspension",
                **base,
                "relation_edges": [
                    base["relation_edges"][0],
                    {
                        **base["relation_edges"][1],
                        "suspension_status": "expired",
                    },
                ],
            },
            {
                "evidence_id": "nd-148-form-03-unofficial-suspension",
                **base,
                "relation_edges": [
                    base["relation_edges"][0],
                    {
                        **base["relation_edges"][1],
                        "official_source_url": "https://example.com/suspension",
                    },
                ],
            },
        ]
    }

    rules, rejections = _validated_form_effectivity_rules(
        payload,
        legal_as_of="2026-07-29",
    )

    assert [item["evidence_id"] for item in rules] == [
        "nd-148-form-03-active"
    ]
    assert rejections == [
        {
            "evidence_id": "nd-148-form-03-no-target-overlap",
            "reason_code": "FORM_EFFECTIVITY_SUSPENDED_REPEAL_SCOPE_UNSAFE",
        },
        {
            "evidence_id": "nd-148-form-03-wrong-suspended-document",
            "reason_code": "FORM_EFFECTIVITY_SUSPENDED_REPEAL_SCOPE_UNSAFE",
        },
        {
            "evidence_id": "nd-148-form-03-inactive-suspension",
            "reason_code": "FORM_EFFECTIVITY_SUSPENDED_REPEAL_SCOPE_UNSAFE",
        },
        {
            "evidence_id": "nd-148-form-03-unofficial-suspension",
            "reason_code": "FORM_EFFECTIVITY_SUSPENDED_REPEAL_SCOPE_UNSAFE",
        },
    ]


def test_effectivity_sidecar_accepts_only_disjoint_expired_form_code_scope() -> None:
    parent_text = "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=167031"
    parent_history = "https://vbpl.vn/TW/Pages/vbpq-lichsu.aspx?ItemID=167031"
    modifier_text = "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=178434"
    package_download = (
        "https://vbpl-bientap-gateway.moj.gov.vn/api/qtdc/public/doc/minio/"
        "buckets/vbpl/167031/VanBanGoc_Th%C3%B4ng%20t%C6%B0-08-2023-"
        "TT-BL%C4%90TBXH.pdf/download"
    )
    base = {
        "issuing_instrument": "08/2023/TT-BLĐTBXH",
        "target_form_code": "18",
        "status": "active",
        "verified_as_of": "2026-07-29",
        "official_source_url": modifier_text,
        "history_source_url": parent_history,
        "form_scope_source_url": parent_text,
        "official_form_package_file_name": (
            "VanBanGoc_Thông tư-08-2023-TT-BLĐTBXH.pdf"
        ),
        "official_form_package_download_url": package_download,
        "official_form_package_sha256": (
            "e341fcdbfd0e10d57fd30bfcf04fdf44f1bd5aa68173deea51e75f6bbf7af7dc"
        ),
        "official_form_package_size": 930044,
        "evidence_basis": "explicit_form_code_scope",
        "relation_edges": [
            {
                "relation": "expires_other_forms_only",
                "modifying_document_number": "09/2025/TT-BNV",
                "effective_from": "2025-07-01",
                "affected_form_codes": [
                    "05",
                    "06",
                    "07",
                    "08",
                    "09",
                    "10",
                    "11",
                    "12",
                    "13",
                    "14",
                    "15",
                ],
                "official_source_url": modifier_text,
            }
        ],
    }
    payload = {
        "form_effectivity": [
            {"evidence_id": "tt-08-form-18-active", **base},
            {
                "evidence_id": "tt-08-form-18-expired",
                **base,
                "relation_edges": [
                    {
                        **base["relation_edges"][0],
                        "affected_form_codes": ["05", "18"],
                    }
                ],
            },
            {
                "evidence_id": "tt-08-form-18-invalid-scope-token",
                **base,
                "relation_edges": [
                    {
                        **base["relation_edges"][0],
                        "affected_form_codes": ["05", "không-rõ"],
                    }
                ],
            },
        ]
    }

    rules, rejections = _validated_form_effectivity_rules(
        payload,
        legal_as_of="2026-07-29",
    )

    assert [item["evidence_id"] for item in rules] == [
        "tt-08-form-18-active"
    ]
    assert rejections == [
        {
            "evidence_id": "tt-08-form-18-expired",
            "reason_code": "FORM_EFFECTIVITY_FORM_SCOPE_UNSAFE",
        },
        {
            "evidence_id": "tt-08-form-18-invalid-scope-token",
            "reason_code": "FORM_EFFECTIVITY_FORM_SCOPE_UNSAFE",
        }
    ]


def test_effectivity_sidecar_understands_nested_provision_scope_and_confirmation() -> None:
    parent_text = "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=148955"
    parent_history = "https://vbpl.vn/tw/Pages/vbpq-lichsu.aspx?ItemID=148955"
    package_download = (
        "https://vbpl-bientap-gateway.moj.gov.vn/api/qtdc/public/doc/minio/"
        "buckets/vbpl/148955/VanBanGoc_20.signed.pdf/download"
    )
    confirmation_source = "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=178293"
    expiry_source = "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=184035"
    base = {
        "issuing_instrument": "20/2021/NĐ-CP",
        "target_form_code": "1A",
        "status": "active",
        "verified_as_of": "2026-07-29",
        "official_source_url": confirmation_source,
        "history_source_url": parent_history,
        "form_scope_source_url": parent_text,
        "target_provisions": ["5.1", "5.2", "5.7", "7.1"],
        "official_form_package_file_name": "VanBanGoc_20.signed.pdf",
        "official_form_package_download_url": package_download,
        "official_form_package_sha256": (
            "627922a72141c70f021d234fb616b42c7418baff391f4f01af3ed8823f469de1"
        ),
        "official_form_package_size": 2552919,
        "evidence_basis": "explicit_partial_provision_scope",
        "relation_edges": [
            {
                "relation": "amends_other_provisions_only",
                "modifying_document_number": "104/2022/NĐ-CP",
                "effective_from": "2023-01-01",
                "affected_provisions": ["8.1.a"],
                "official_source_url": (
                    "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=167220"
                ),
            },
            {
                "relation": "confirms_target_form_binding",
                "modifying_document_number": "147/2025/NĐ-CP",
                "effective_from": "2025-07-01",
                "confirmed_provisions": ["7.1"],
                "official_source_url": confirmation_source,
            },
            {
                "relation": "expires_other_provisions_only",
                "modifying_document_number": "176/2025/NĐ-CP",
                "effective_from": "2025-07-01",
                "affected_provisions": ["5.5.b", "5.5.c"],
                "official_source_url": expiry_source,
            },
        ],
    }
    payload = {
        "form_effectivity": [
            {"evidence_id": "nd-20-form-1a-active", **base},
            {
                "evidence_id": "nd-20-form-1d-nested-overlap",
                **base,
                "target_form_code": "1D",
                "target_provisions": ["5.5", "7.1"],
            },
            {
                "evidence_id": "nd-20-form-1a-unconfirmed-binding",
                **base,
                "relation_edges": [
                    {
                        **base["relation_edges"][1],
                        "confirmed_provisions": ["27.2.a"],
                    }
                ],
            },
        ]
    }

    rules, rejections = _validated_form_effectivity_rules(
        payload,
        legal_as_of="2026-07-29",
    )

    assert [item["evidence_id"] for item in rules] == ["nd-20-form-1a-active"]
    assert rejections == [
        {
            "evidence_id": "nd-20-form-1d-nested-overlap",
            "reason_code": "FORM_EFFECTIVITY_PROVISION_SCOPE_UNSAFE",
        },
        {
            "evidence_id": "nd-20-form-1a-unconfirmed-binding",
            "reason_code": "FORM_EFFECTIVITY_PROVISION_SCOPE_UNSAFE",
        },
    ]


def test_superseded_appendix_is_never_eligible() -> None:
    decision = classify_form_effectivity(
        {
            "effStatus": "Hết hiệu lực một phần",
            "effFrom": "2020-01-01",
            "appendix_effectivity": [
                {
                    "appendix_identifier": "I",
                    "status": "superseded",
                    "effective_to": "2025-01-01",
                    "replacement_document_number": "02/2025/TT-BTP",
                    "official_source_url": "https://vbpl.vn/van-ban/replacement",
                    "verified_as_of": "2026-07-27",
                }
            ],
        },
        legal_as_of="2026-07-27",
        appendix_identifier="I",
    )

    assert decision["eligible"] is False
    assert decision["reason_code"] == "FORM_APPENDIX_SUPERSEDED"
    assert decision["replacement_document_number"] == "02/2025/TT-BTP"
    assert classify_effectivity(
        {"effStatus": "Hết hiệu lực", "effFrom": "2021-07-01", "effTo": "2026-07-01"},
        legal_as_of="2026-07-27",
    )["reason_code"] == "ISSUING_INSTRUMENT_EXPIRED"
    assert classify_effectivity(
        {"effStatus": "Hết hiệu lực một phần", "effFrom": "2025-07-01"},
        legal_as_of="2026-07-27",
    )["reason_code"] == "PARTIAL_EFFECTIVITY_REQUIRES_REVIEW"
    assert classify_effectivity(
        {
            "effStatus": {"code": "CHL", "name": "Còn hiệu lực"},
            "effFrom": "2025-07-01T00:00:00",
        },
        legal_as_of="2026-07-27",
    )["eligible"] is True


def test_select_standalone_file_requires_exact_form_code() -> None:
    files = [
        {"fileName": "Thông tư 53.pdf", "relatedType": 1},
        {"fileName": "1. Mẫu CT01 ban hành kèm theo Thông tư 53.doc", "relatedType": 2},
        {"fileName": "2. Mẫu CT02 ban hành kèm theo Thông tư 53.doc", "relatedType": 2},
    ]

    selected = select_standalone_form_file(files, "CT01")

    assert selected is not None
    assert "CT01" in selected["fileName"]
    assert select_standalone_form_file(files, "CT03") is None


def test_select_standalone_file_rejects_ambiguous_matches() -> None:
    files = [
        {"fileName": "Mẫu số 01.doc", "relatedType": 2},
        {"fileName": "Phụ lục Mẫu số 01.pdf", "relatedType": 2},
    ]

    assert select_standalone_form_file(files, "01") is None


def test_pdf_package_range_stops_before_next_form() -> None:
    pages = [
        "Phụ lục danh mục biểu mẫu",
        "Mẫu số 01\nTỜ KHAI\nTrang thứ nhất",
        "Nội dung tiếp theo của tờ khai",
        "Mẫu số 02\nĐƠN ĐỀ NGHỊ",
    ]

    assert locate_pdf_form_page_range(pages, "01") == (1, 3)
    assert locate_pdf_form_page_range(pages, "03") is None
    assert locate_pdf_form_page_range(["Mẫu số 01 và Mẫu số 02"], "01") is None


def test_pdf_package_range_prefers_exact_form_header_over_earlier_legal_reference() -> None:
    pages = [
        "Điều 2. Hồ sơ gồm Giấy đề nghị theo Mẫu số 01.",
        "Mẫu số 01\nBan hành kèm theo Thông tư 01/2025/TT-BTP\nTỜ KHAI",
        "Nội dung tiếp theo của Mẫu số 01",
        "Mẫu số 02\nBan hành kèm theo Thông tư 01/2025/TT-BTP\nĐƠN ĐỀ NGHỊ",
    ]

    assert locate_pdf_form_page_range(pages, "01") == (1, 3)


def test_pdf_package_range_rejects_citation_only_page_with_single_form_code() -> None:
    pages = [
        (
            "Điều 35. Trình tự, thủ tục giải thể trung tâm\n"
            "Tờ trình đề nghị giải thể theo Mẫu số 06 Phụ lục II kèm theo "
            "Nghị định này."
        )
    ]

    assert locate_pdf_form_page_range(pages, "06", appendix_identifier="II") is None


def test_pdf_package_range_accepts_scanned_header_without_attachment_caption() -> None:
    pages = [
        "Điều 11. Người đề nghị lập tờ khai theo Mẫu số 04.",
        "Mẫu số 03\nTỜ KHAI CÁ NHÂN",
        "Nội dung tiếp theo của Mẫu số 03",
        "Mẫu số 04\nCỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nTỜ KHAI",
        "Nội dung tiếp theo của tờ khai",
        "14\n\nMẫu số 05\nCỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nBẢN KHAI",
    ]

    assert locate_pdf_form_page_range(pages, "04") == (3, 5)
    assert locate_pdf_form_page_range(pages, "05") == (5, 6)


def test_pdf_package_range_accepts_exact_header_after_short_section_intro() -> None:
    pages = [
        "Phần 1. Hồ sơ sử dụng Mẫu số 01 theo quy định.",
        (
            "Phần 6. Các biểu mẫu kèm theo\n"
            "A. Biểu mẫu về tiếp tục lưu hành thiết bị y tế\n"
            "áp dụng cho thiết bị loại A, B\n"
            "Mẫu số 01\n"
            "VĂN BẢN ĐỀ NGHỊ TIẾP TỤC LƯU HÀNH"
        ),
        "Nội dung tiếp theo của Mẫu số 01",
        "Mẫu số 02\nVĂN BẢN ĐỒNG Ý TIẾP TỤC LƯU HÀNH",
    ]

    assert locate_pdf_form_page_range(pages, "01") == (1, 3)


def test_pdf_package_range_accepts_exact_miu_so_ocr_heading() -> None:
    pages = [
        "Mẫu số 11\nQUYẾT ĐỊNH CHỈ ĐỊNH CƠ SỞ KIỂM NGHIỆM",
        "Danh mục chỉ tiêu kèm theo Mẫu số 11",
        "31\nMiu so 12\nĐơn đề nghị cấp giấy chứng nhận đối với thực phẩm xuất khẩu",
        "Nội dung tiếp theo của đơn đề nghị",
        "PHỤ LỤC VI\nHỒ SƠ THỤ TINH TRONG ỐNG NGHIỆM",
    ]

    assert locate_pdf_form_page_range(pages, "12") == (2, 4)


def test_pdf_package_range_accepts_fragmented_glyphs_only_in_exact_heading() -> None:
    pages = [
        "Danh mục: Mẫu số 18; Mẫu số 19; Mẫu số 20",
        "28\nM\nẫ\nu s\nố\n18\nCỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nTỜ KHAI",
        "29\nM\nẫ\nu s\nố\n19\nCỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nTỜ KHAI",
        "Ghi chú áp dụng chung cho Mẫu số 18 và Mẫu số 19.",
        "31\nMẫu số 20\nCỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nTỜ KHAI",
    ]

    assert locate_pdf_form_page_range(pages, "18") == (1, 2)
    assert locate_pdf_form_page_range(pages, "19") == (2, 4)


def test_pdf_package_range_uses_exact_ocr_header_despite_spurious_body_code() -> None:
    pages = [
        "Mẫu số 1c\nTỜ KHAI\nGiấy M6I bị OCR sai trong phần nội dung",
        "Mẫu so 1đ\nTỜ KHAI ĐỀ NGHỊ TRỢ GIÚP XÃ HỘI",
        "Mẫu số 2a\nTỜ KHAI HỘ GIA ĐÌNH",
    ]

    assert locate_pdf_form_page_range(pages, "1C") == (0, 1)
    assert locate_pdf_form_page_range(pages, "1Đ") == (1, 2)


def test_pdf_package_range_disambiguates_repeated_code_by_appendix() -> None:
    pages = [
        "PHỤ LỤC I\nDANH MỤC BIỂU MẪU",
        "Mẫu số 01\nBẢN CÔNG BỐ CƠ SỞ HƯỚNG DẪN THỰC HÀNH",
        "Nội dung Mẫu 01 thuộc Phụ lục I",
        "PHỤ LỤC II\nDANH MỤC BIỂU MẪU",
        "Mẫu số 01\nDANH SÁCH ĐĂNG KÝ HÀNH NGHỀ",
        "Nội dung Mẫu 01 thuộc Phụ lục II",
    ]

    assert locate_pdf_form_page_range(pages, "01") is None
    assert locate_pdf_form_page_range(
        pages, "01", appendix_identifier="I"
    ) == (1, 3)
    assert locate_pdf_form_page_range(
        pages, "01", appendix_identifier="II"
    ) == (4, 6)


def test_pdf_package_range_stops_at_ocr_phy_luc_heading() -> None:
    pages = [
        "Phụ lục IV\nCÁC BIỂU MẪU KHÁM BỆNH NHÂN ĐẠO",
        "Mẫu 03 - Kế hoạch tổ chức khám bệnh, chữa bệnh nhân đạo",
        "Nội dung tiếp theo của kế hoạch",
        "Phy luc V\nHƯỚNG DẪN XẾP CẤP CHUYÊN MÔN KỸ THUẬT",
    ]

    assert locate_pdf_form_page_range(
        pages, "03", appendix_identifier="IV"
    ) == (1, 3)


def test_pending_candidate_is_never_auto_approved() -> None:
    record = build_pending_candidate(
        group={
            "group_id": "group-1",
            "form_code": "CT01",
            "form_name": "Mẫu CT01",
            "issuing_instrument": "53/2025/TT-BCA",
            "procedure_ids": ["1.001"],
            "domain": "cu_tru_an_ninh",
            "source_tier": "central",
        },
        document={
            "id": 179354,
            "docNum": "53/2025/TT-BCA",
            "title": "Thông tư số 53/2025/TT-BCA",
            "effFrom": "2025-07-01",
            "effTo": None,
            "effStatus": "Còn hiệu lực",
        },
        artifact={
            "file_name": "Mẫu CT01.doc",
            "local_path": "data/uploads/three_tier_form_candidates/ct01.doc",
            "download_url": "https://vbpl-bientap-gateway.moj.gov.vn/api/file",
            "sha256": "a" * 64,
            "size_bytes": 128,
        },
        legal_as_of="2026-07-27",
    )

    assert record["procedure_id"] == "1.001"
    assert record["approved"] is False
    assert record["runtime_eligible"] is False
    assert record["legal_review_status"] == "candidate_pending_review"
    assert record["hard_gate_reason_codes"] == ["HUMAN_LEGAL_REVIEW_REQUIRED"]
    assert record["provenance"]["document_id"] == "179354"
    assert record["provenance"]["retrieved_at"]


def test_pending_candidate_uses_exact_appendix_effectivity_for_partial_document() -> None:
    record = build_pending_candidate(
        group={
            "group_id": "group-partial",
            "form_code": "04",
            "form_name": "Mẫu số 04",
            "issuing_instrument": "09/2025/TT-BNV",
            "appendix_identifier": "I",
            "procedure_ids": ["1.013749"],
            "domain": "an_sinh_y_te_giao_duc",
            "source_tier": "central",
        },
        document={
            "id": 200001,
            "docNum": "09/2025/TT-BNV",
            "effFrom": "2025-07-01",
            "effTo": None,
            "effStatus": "Hết hiệu lực một phần",
            "appendix_effectivity": [
                {
                    "appendix_identifier": "I",
                    "status": "active",
                    "official_source_url": "https://vbpl.vn/van-ban/example",
                    "verified_as_of": "2026-07-29",
                }
            ],
        },
        artifact={
            "file_name": "Mẫu số 04.doc",
            "local_path": "data/uploads/three_tier_form_candidates/mau-04.doc",
            "download_url": "https://vbpl-bientap-gateway.moj.gov.vn/api/file",
            "sha256": "b" * 64,
            "size_bytes": 128,
        },
        legal_as_of="2026-07-29",
    )

    assert record["effectivity_reason_code"] == "APPENDIX_EFFECTIVITY_VERIFIED"
    assert record["effectivity_source_url"] == "https://vbpl.vn/van-ban/example"
    assert record["approved"] is False
    assert record["runtime_eligible"] is False


def test_pending_candidate_preserves_dvc_provenance_without_auto_approval() -> None:
    procedure_url = (
        "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/"
        "019d2bff-2d2f-73b6-a226-b9bad3aca884"
    )
    attachment_api = (
        "https://dichvucong.gov.vn/api/v1/submitting/preview-attachment"
    )
    record = build_pending_candidate(
        group={
            "group_id": "group-dvc-07",
            "form_code": "07",
            "form_name": "Mẫu số 07",
            "issuing_instrument": "91/2016/NĐ-CP",
            "appendix_identifier": "I",
            "procedure_ids": ["1.013870"],
            "domain": "an_sinh_y_te_giao_duc",
            "source_tier": "central",
            "executing_level": "province",
        },
        document={
            "id": "112021",
            "docNum": "91/2016/NĐ-CP",
            "detailUrl": "https://vbpl.vn/van-ban/chi-tiet/--112021",
            "effFrom": "2016-07-01",
            "effTo": None,
            "effStatus": "Hết hiệu lực một phần",
            "appendix_effectivity": [
                {
                    "appendix_identifier": "I",
                    "target_form_code": "07",
                    "status": "current",
                    "official_source_url": procedure_url,
                    "verified_as_of": "2026-07-29",
                }
            ],
        },
        artifact={
            "file_name": "form-07.docx",
            "local_path": "data/uploads/forms/form-07.docx",
            "download_url": procedure_url,
            "sha256": "c" * 64,
            "size_bytes": 1024,
            "source_page_url": procedure_url,
            "source_retrieval_url": attachment_api,
            "source_retrieval_method": "POST",
            "source_attachment_id": "019eb4af-a3af-7143-b664-6dd27b38b9de",
            "source_file_name": "QĐ 1684 BYT.docx",
            "source_package_sha256": "d" * 64,
            "source_package_size_bytes": 170394,
            "publication_decision_number": "1684/QĐ-BYT",
            "extraction": {
                "kind": "structural_docx_form_boundary",
                "complete": True,
                "element_range": [12, 31],
            },
            "provenance_kind": "official_dvc_attachment",
            "publisher": "Cổng Dịch vụ công quốc gia",
        },
        legal_as_of="2026-07-29",
    )

    assert record["source_url"] == procedure_url
    assert record["source_download_url"] == procedure_url
    assert record["publisher"] == "Cổng Dịch vụ công quốc gia"
    assert record["provenance"]["kind"] == "official_dvc_attachment"
    assert record["provenance"]["source_retrieval_url"] == attachment_api
    assert (
        record["provenance"]["source_attachment_id"]
        == "019eb4af-a3af-7143-b664-6dd27b38b9de"
    )
    assert record["administrative_level"] == "province"
    assert record["source_attachment_id"] == "019eb4af-a3af-7143-b664-6dd27b38b9de"
    assert record["source_package_sha256"] == "d" * 64
    assert record["source_package_size_bytes"] == 170394
    assert record["provenance_kind"] == "official_dvc_attachment"
    assert record["extraction"]["kind"] == "structural_docx_form_boundary"
    assert record["approved"] is False
    assert record["runtime_eligible"] is False
    assert record["hard_gate_reason_codes"] == [
        "HUMAN_LEGAL_REVIEW_REQUIRED"
    ]


def test_pending_candidate_merge_is_idempotent_and_detects_checksum_drift() -> None:
    original = {
        "id": "three-tier-1",
        "sha256": "a" * 64,
        "approved": False,
        "runtime_eligible": False,
    }

    first = merge_pending_records([original], [dict(original)])
    assert first["action_counts"] == {"created": 0, "updated": 0, "unchanged": 1}
    assert first["records"] == [original]

    changed = {**original, "sha256": "b" * 64}
    try:
        merge_pending_records([original], [changed])
    except ValueError as exc:
        assert str(exc) == "candidate_checksum_drift"
    else:
        raise AssertionError("checksum drift must fail closed")


def test_pending_candidate_merge_retains_existing_human_disposition() -> None:
    existing = {
        "id": "three-tier-reviewed",
        "sha256": "a" * 64,
        "source_package_sha256": "b" * 64,
        "official_procedure_code": "1.001776",
        "procedure_id": "1.001776",
        "proposed_canonical_form_id": "form-three-tier-reviewed",
        "review_status": "approved",
        "legal_review_status": "catalog_sync_required",
        "catalog_status": "legal_review_required",
        "reviewed_by": "user_account:admin",
        "reviewed_at": "2026-07-30T14:00:00+00:00",
        "is_approved": True,
        "approved": False,
        "runtime_eligible": False,
        "hard_gate_reason_codes": ["HUMAN_LEGAL_REVIEW_REQUIRED"],
        "local_path": "data/uploads/forms/original.pdf",
    }
    incoming = {
        "id": "three-tier-reviewed",
        "sha256": "a" * 64,
        "source_package_sha256": "b" * 64,
        "official_procedure_code": "1.001776",
        "procedure_id": "1.001776",
        "proposed_canonical_form_id": "form-three-tier-reviewed",
        "review_status": "candidate_pending_review",
        "legal_review_status": "candidate_pending_review",
        "catalog_status": "candidate_pending_review",
        "is_approved": False,
        "approved": False,
        "runtime_eligible": False,
        "hard_gate_reason_codes": ["HUMAN_LEGAL_REVIEW_REQUIRED"],
        "local_path": "data/uploads/forms/retrieved.pdf",
        "provenance": {"retrieved_at": "2026-07-30T15:00:00+00:00"},
    }

    result = merge_pending_records([existing], [incoming])

    record = result["records"][0]
    assert result["action_counts"] == {"created": 0, "updated": 1, "unchanged": 0}
    assert result["preserved_review_dispositions"] == 1
    assert record["review_status"] == "approved"
    assert record["legal_review_status"] == "catalog_sync_required"
    assert record["catalog_status"] == "legal_review_required"
    assert record["reviewed_by"] == "user_account:admin"
    assert record["is_approved"] is True
    assert record["approved"] is False
    assert record["runtime_eligible"] is False
    assert record["local_path"] == "data/uploads/forms/original.pdf"
    assert record["provenance"]["retrieved_at"] == "2026-07-30T15:00:00+00:00"


def test_stale_generated_candidate_is_quarantined_not_deleted() -> None:
    records = [
        {
            "id": "three-tier-old",
            "approved": False,
            "runtime_eligible": False,
            "hard_gate_reason_codes": ["HUMAN_LEGAL_REVIEW_REQUIRED"],
        },
        {"id": "manual-record", "approved": False},
    ]

    result = quarantine_stale_pending_records(records, active_ids=set())

    assert len(result["records"]) == 2
    assert result["quarantined_count"] == 1
    assert result["records"][0]["is_quarantined"] is True
    assert result["records"][0]["runtime_eligible"] is False
    assert result["records"][0]["hard_gate_reason_codes"] == [
        "SOURCE_RESOLUTION_STALE"
    ]
    assert result["records"][1] == records[1]


def test_scanned_pdf_package_uses_only_complete_page_ocr(
    tmp_path, monkeypatch
) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    payload = BytesIO()
    writer.write(payload)

    monkeypatch.setattr(
        "api.crawlers.ocr_extractor.extract_ocr_from_pdf_bytes",
        lambda _content, **_kwargs: {
            "complete": True,
            "text": (
                "--- Trang 1 ---\nMẫu số 01\nTỜ KHAI\nNội dung biểu mẫu\n"
                "--- Trang 2 ---\nMẫu số 02\nĐƠN KHÁC"
            ),
        },
    )

    artifact = _extract_from_pdf_package(
        payload.getvalue(),
        form_code="01",
        download_url="https://vbpl.vn/package.pdf",
        download_dir=tmp_path,
        ocr_cache_dir=tmp_path / "ocr-cache",
    )

    assert artifact is not None
    assert artifact["source_pages_zero_based"] == [0]
    assert artifact["source_package_sha256"]


def test_partial_ocr_never_extracts_a_scanned_form(
    tmp_path, monkeypatch
) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    payload = BytesIO()
    writer.write(payload)
    monkeypatch.setattr(
        "api.crawlers.ocr_extractor.extract_ocr_from_pdf_bytes",
        lambda _content, **_kwargs: {
            "complete": False,
            "text": "--- Trang 1 ---\nMẫu số 01",
        },
    )

    assert (
        _extract_from_pdf_package(
            payload.getvalue(),
            form_code="01",
            download_url="https://vbpl.vn/package.pdf",
            download_dir=tmp_path,
            ocr_cache_dir=tmp_path / "ocr-cache",
        )
        is None
    )


def test_scanned_pdf_can_defer_ocr_until_machine_readable_packages_are_tried(
    tmp_path, monkeypatch
) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    payload = BytesIO()
    writer.write(payload)
    monkeypatch.setattr(
        "api.crawlers.ocr_extractor.extract_ocr_from_pdf_bytes",
        lambda _content, **_kwargs: (_ for _ in ()).throw(
            AssertionError("OCR must be deferred")
        ),
    )

    assert (
        _extract_from_pdf_package(
            payload.getvalue(),
            form_code="01",
            download_url="https://vbpl.vn/package.pdf",
            download_dir=tmp_path,
            allow_ocr=False,
            ocr_cache_dir=tmp_path / "ocr-cache",
        )
        is None
    )


def test_legacy_doc_package_preserves_original_checksum(
    tmp_path, monkeypatch
) -> None:
    document = Document()
    document.add_paragraph("Mẫu số 01")
    document.add_paragraph("TỜ KHAI")
    document.add_paragraph("Nội dung " * 40)
    document.add_paragraph("Mẫu số 02")
    output = BytesIO()
    document.save(output)
    source = bytes.fromhex("D0CF11E0A1B11AE1") + b"legacy-package"

    monkeypatch.setattr(
        "api.crawlers.office_converter.convert_legacy_doc_bytes",
        lambda *_args, **_kwargs: {
            "status": "ok",
            "complete": True,
            "content": output.getvalue(),
            "sha256": "c" * 64,
        },
    )

    artifact, reason = _extract_from_legacy_doc_package(
        source,
        filename="package.doc",
        form_code="01",
        download_url="https://vbpl.vn/package.doc",
        download_dir=tmp_path,
    )

    assert reason == "EXTRACTED_FROM_LEGACY_OFFICIAL_DOC_PACKAGE"
    assert artifact is not None
    assert artifact["source_file_format"] == "doc"
    assert artifact["conversion"]["complete"] is True
    assert artifact["source_package_sha256"] != artifact["sha256"]


def test_docx_extraction_uses_appendix_header_not_earlier_legal_reference(
    tmp_path,
) -> None:
    document = Document()
    document.add_paragraph(
        "Hồ sơ gồm đơn đề nghị theo Mẫu số 02 tại Phụ lục ban hành kèm theo."
    )
    document.add_paragraph("Điều tiếp theo")
    document.add_paragraph("Mẫu số 02")
    document.add_paragraph("ĐƠN ĐỀ NGHỊ CẤP GIẤY CHỨNG NHẬN")
    document.add_paragraph("Nội dung biểu mẫu " * 40)
    document.add_paragraph("Mẫu số 03")
    output = BytesIO()
    document.save(output)

    artifact = _extract_from_docx_package(
        output.getvalue(),
        form_code="02",
        download_url="https://vbpl.vn/package.docx",
        download_dir=tmp_path,
    )

    assert artifact is not None
    extracted = Document(tmp_path / artifact["file_name"])
    extracted_text = "\n".join(
        paragraph.text for paragraph in extracted.paragraphs
    )
    assert "ĐƠN ĐỀ NGHỊ CẤP GIẤY CHỨNG NHẬN" in extracted_text
    assert "Hồ sơ gồm đơn đề nghị" not in extracted_text


def test_docx_extraction_accepts_form_header_before_appendix_and_stops_at_next_procedure(
    tmp_path,
) -> None:
    document = Document()
    document.add_paragraph("Mẫu số 07")
    document.add_paragraph("PHỤ LỤC I")
    document.add_paragraph("VĂN BẢN ĐỀ NGHỊ CẤP LẠI")
    document.add_paragraph("Nội dung biểu mẫu " * 40)
    next_procedure = document.add_table(rows=1, cols=1)
    next_procedure.cell(0, 0).text = (
        "Mã thủ tục hành chính: 1.013872 "
        "Thông báo thay đổi nội dung, hình thức nhãn chế phẩm"
    )
    document.add_paragraph("Mẫu số 09")
    document.add_paragraph("Nội dung của thủ tục tiếp theo " * 20)
    output = BytesIO()
    document.save(output)

    artifact = _extract_from_docx_package(
        output.getvalue(),
        form_code="07",
        appendix_identifier="I",
        download_url="https://dichvucong.gov.vn/quyet-dinh-cong-bo/example",
        download_dir=tmp_path,
    )

    assert artifact is not None
    extracted = Document(tmp_path / artifact["file_name"])
    extracted_text = " ".join(
        str(node.text or "")
        for node in extracted.element.body.iter()
        if str(node.tag).endswith("}t")
    )
    assert "Mẫu số 07" in extracted_text
    assert "PHỤ LỤC I" in extracted_text
    assert "VĂN BẢN ĐỀ NGHỊ CẤP LẠI" in extracted_text
    assert "1.013872" not in extracted_text
    assert "Mẫu số 09" not in extracted_text


def test_docx_extraction_checksum_is_stable_across_repeat_runs() -> None:
    document = Document()
    document.add_paragraph("stable source package")
    source = BytesIO()
    document.save(source)

    first_document = Document(BytesIO(source.getvalue()))
    first_document.core_properties.modified = datetime(
        2026, 7, 29, 12, 0, tzinfo=timezone.utc
    )
    first_output = BytesIO()
    first_document.save(first_output)

    second_document = Document(BytesIO(source.getvalue()))
    second_document.core_properties.modified = datetime(
        2026, 7, 29, 12, 1, tzinfo=timezone.utc
    )
    second_output = BytesIO()
    second_document.save(second_output)

    assert first_output.getvalue() != second_output.getvalue()
    first_canonical = _deterministic_docx_bytes(first_output.getvalue())
    second_canonical = _deterministic_docx_bytes(second_output.getvalue())

    assert first_canonical == second_canonical
    assert _deterministic_docx_bytes(first_canonical) == first_canonical
    assert Document(BytesIO(first_canonical)).paragraphs[0].text == (
        "stable source package"
    )


def test_docx_extraction_disambiguates_repeated_code_by_appendix(tmp_path) -> None:
    document = Document()
    document.add_paragraph("PHỤ LỤC I")
    document.add_paragraph("Mẫu số 01")
    document.add_paragraph("BẢN CÔNG BỐ CƠ SỞ HƯỚNG DẪN THỰC HÀNH " * 8)
    document.add_paragraph("PHỤ LỤC II")
    document.add_paragraph("Mẫu số 01")
    document.add_paragraph("DANH SÁCH ĐĂNG KÝ HÀNH NGHỀ " * 8)
    document.add_paragraph("Mẫu số 02")
    output = BytesIO()
    document.save(output)

    assert _extract_from_docx_package(
        output.getvalue(),
        form_code="01",
        download_url="https://vbpl.vn/package.docx",
        download_dir=tmp_path,
    ) is None

    artifact = _extract_from_docx_package(
        output.getvalue(),
        form_code="01",
        appendix_identifier="II",
        download_url="https://vbpl.vn/package.docx",
        download_dir=tmp_path,
    )

    assert artifact is not None
    extracted = Document(tmp_path / artifact["file_name"])
    extracted_text = "\n".join(
        paragraph.text for paragraph in extracted.paragraphs
    )
    assert "DANH SÁCH ĐĂNG KÝ HÀNH NGHỀ" in extracted_text
    assert "BẢN CÔNG BỐ CƠ SỞ HƯỚNG DẪN THỰC HÀNH" not in extracted_text


def test_docx_extraction_reconstructs_form_heading_split_across_runs(
    tmp_path,
) -> None:
    document = Document()
    document.add_paragraph("PHỤ LỤC I")
    heading = document.add_paragraph()
    for fragment in ("M", "ẫ", "u ", "s", "ố ", "01"):
        heading.add_run(fragment)
    document.add_paragraph("ĐƠN ĐỀ NGHỊ GIA HẠN THỜI HẠN SỞ HỮU NHÀ Ở")
    document.add_paragraph("Nội dung biểu mẫu " * 40)
    document.add_paragraph("Mẫu số 02")
    output = BytesIO()
    document.save(output)

    artifact = _extract_from_docx_package(
        output.getvalue(),
        form_code="01",
        appendix_identifier="I",
        download_url="https://vbpl.vn/package.docx",
        download_dir=tmp_path,
    )

    assert artifact is not None
    extracted = Document(tmp_path / artifact["file_name"])
    assert "Mẫu số 01" in "\n".join(
        paragraph.text for paragraph in extracted.paragraphs
    )


def test_official_attachment_cache_reuses_validated_bytes(tmp_path) -> None:
    class Response:
        content = b"%PDF-1.7 cached\n" + (b"0" * 256)
        headers = {"content-type": "application/pdf"}

        def raise_for_status(self):
            return None

    class Client:
        calls = 0

        def get(self, _url):
            self.calls += 1
            return Response()

    client = Client()
    file_record = {"fileName": "package.pdf"}

    # validate_download and malware checks are exercised by existing crawler
    # tests; this test focuses on the idempotent cache boundary.
    first = _validated_content(
        client,
        document_id="doc-1",
        file_record=file_record,
        cache_dir=tmp_path,
    )
    second = _validated_content(
        client,
        document_id="doc-1",
        file_record=file_record,
        cache_dir=tmp_path,
    )

    assert first == second
    assert client.calls == 1


def test_dvc_attachment_download_posts_exact_file_id_and_reuses_safe_cache(
    tmp_path,
) -> None:
    document = Document()
    document.add_paragraph("Mẫu số 05")
    document.add_paragraph("PHỤ LỤC I")
    document.add_paragraph("Nội dung biểu mẫu " * 40)
    output = BytesIO()
    document.save(output)
    content = output.getvalue()
    endpoint = (
        "https://dichvucong.gov.vn/api/v1/submitting/preview-attachment"
    )
    procedure_url = (
        "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/"
        "019d2bff-2d36-7533-81f7-2f3380a96f1b"
    )
    request = {
        "api_url": endpoint,
        "attachment_id": "019e2aed-9d64-73ee-a43e-245d6a2a06e6",
        "file_name": "maudon.docx",
        "sha256": hashlib.sha256(content).hexdigest(),
        "size_bytes": len(content),
        "referer_url": procedure_url,
        "source_page_url": procedure_url,
    }

    class Response:
        url = httpx.URL(endpoint)
        headers = {
            "content-type": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        }

        def __init__(self):
            self.content = content

        def raise_for_status(self):
            return None

    class Client:
        def __init__(self):
            self.calls = 0

        def post(self, url, *, json, headers):
            self.calls += 1
            assert url == endpoint
            assert json == {
                "fileId": "019e2aed-9d64-73ee-a43e-245d6a2a06e6"
            }
            assert headers["Origin"] == "https://dichvucong.gov.vn"
            assert headers["Referer"] == procedure_url
            return Response()

    client = Client()
    first = _validated_dvc_attachment_content(
        client,
        request=request,
        cache_dir=tmp_path,
    )
    second = _validated_dvc_attachment_content(
        client,
        request=request,
        cache_dir=tmp_path,
    )

    assert first == second == (content, endpoint, "docx")
    assert client.calls == 1


def test_dvc_attachment_ignores_corrupt_cache_metadata_and_revalidates(
    tmp_path,
) -> None:
    document = Document()
    document.add_paragraph("Máº«u sá»‘ 05")
    document.add_paragraph("PHá»¤ Lá»¤C I")
    output = BytesIO()
    document.save(output)
    content = output.getvalue()
    endpoint = (
        "https://dichvucong.gov.vn/api/v1/submitting/preview-attachment"
    )
    procedure_url = (
        "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/"
        "019d2bff-2d36-7533-81f7-2f3380a96f1b"
    )
    request = {
        "api_url": endpoint,
        "attachment_id": "019e2aed-9d64-73ee-a43e-245d6a2a06e6",
        "file_name": "maudon.docx",
        "sha256": hashlib.sha256(content).hexdigest(),
        "size_bytes": len(content),
        "referer_url": procedure_url,
        "source_page_url": procedure_url,
    }

    class Response:
        url = httpx.URL(endpoint)
        headers = {
            "content-type": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        }

        def __init__(self):
            self.content = content

        def raise_for_status(self):
            return None

    class Client:
        def __init__(self):
            self.calls = 0

        def post(self, *_args, **_kwargs):
            self.calls += 1
            return Response()

    first_client = Client()
    _validated_dvc_attachment_content(
        first_client,
        request=request,
        cache_dir=tmp_path,
    )
    cache_metadata = next(tmp_path.glob("*.json"))
    cache_metadata.write_text("{not-json", encoding="utf-8")

    second_client = Client()
    result = _validated_dvc_attachment_content(
        second_client,
        request=request,
        cache_dir=tmp_path,
    )

    metadata = json.loads(cache_metadata.read_text(encoding="utf-8"))
    metadata["file_format"] = "txt"
    cache_metadata.write_text(json.dumps(metadata), encoding="utf-8")
    third_client = Client()
    third = _validated_dvc_attachment_content(
        third_client,
        request=request,
        cache_dir=tmp_path,
    )

    assert result == (content, endpoint, "docx")
    assert third == (content, endpoint, "docx")
    assert first_client.calls == 1
    assert second_client.calls == 1
    assert third_client.calls == 1


def test_dvc_attachment_download_rejects_checksum_drift_without_cache(
    tmp_path,
) -> None:
    content = b"PK\x03\x04" + (b"0" * 512)
    endpoint = (
        "https://dichvucong.gov.vn/api/v1/submitting/preview-attachment"
    )
    request = {
        "api_url": endpoint,
        "attachment_id": "019e2aed-9d64-73ee-a43e-245d6a2a06e6",
        "file_name": "maudon.docx",
        "sha256": "a" * 64,
        "size_bytes": len(content),
        "referer_url": (
            "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/"
            "019d2bff-2d36-7533-81f7-2f3380a96f1b"
        ),
        "source_page_url": (
            "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/"
            "019d2bff-2d36-7533-81f7-2f3380a96f1b"
        ),
    }

    class Response:
        url = httpx.URL(endpoint)
        headers = {
            "content-type": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        }

        def __init__(self):
            self.content = content

        def raise_for_status(self):
            return None

    class Client:
        def post(self, *_args, **_kwargs):
            return Response()

    with pytest.raises(ValueError, match="OFFICIAL_FORM_CHECKSUM_DRIFT"):
        _validated_dvc_attachment_content(
            Client(),
            request=request,
            cache_dir=tmp_path,
        )
    assert list(tmp_path.iterdir()) == []


def test_group_artifact_prefers_checksum_bound_dvc_attachment(
    tmp_path, monkeypatch
) -> None:
    class Client:
        def __init__(self, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    request = {
        "api_url": (
            "https://dichvucong.gov.vn/api/v1/submitting/"
            "preview-attachment"
        ),
        "attachment_id": "019e2aed-9d64-73ee-a43e-245d6a2a06e6",
        "file_name": "maudon.docx",
        "sha256": "d" * 64,
        "size_bytes": 24924,
        "referer_url": "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/test",
        "source_page_url": (
            "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/test"
        ),
    }
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.httpx.Client", Client
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_dvc_attachment_content",
        lambda *_args, **_kwargs: (
            b"PK\x03\x04 current dvc",
            request["api_url"],
            "docx",
        ),
    )
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_content",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("VBPL fallback must not run")
        ),
    )

    def extract_docx(content, **kwargs):
        assert content == b"PK\x03\x04 current dvc"
        assert kwargs["form_code"] == "05"
        assert kwargs["appendix_identifier"] == "I"
        assert kwargs["download_url"] == request["source_page_url"]
        assert kwargs["source_package_sha256"] == hashlib.sha256(
            content
        ).hexdigest()
        return {
            "file_name": "form.docx",
            "local_path": "data/uploads/forms/form.docx",
            "download_url": kwargs["download_url"],
            "sha256": "e" * 64,
            "size_bytes": 1024,
            "file_format": "docx",
        }

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._extract_from_docx_package",
        extract_docx,
    )

    artifact, reason = resolve_group_artifact(
        {
            "group_id": "group-current-dvc",
            "form_code": "05",
            "appendix_identifier": "I",
            "official_dvc_attachment_required": True,
            "official_dvc_attachment_reason_code": None,
            "official_dvc_attachment_requests": [request],
        },
        {
            "document": {"id": "112021"},
            "files": [{"fileName": "old-package.pdf"}],
        },
        timeout=1,
        download_dir=tmp_path,
    )

    assert reason == "EXTRACTED_FROM_OFFICIAL_DVC_ATTACHMENT"
    assert artifact is not None
    assert artifact["provenance_kind"] == "official_dvc_attachment"
    assert artifact["publisher"] == "Cổng Dịch vụ công quốc gia"
    assert artifact["source_attachment_id"] == request["attachment_id"]
    assert artifact["source_retrieval_url"] == request["api_url"]
    assert artifact["source_page_url"] == request["source_page_url"]


def test_direct_official_package_rejects_unallowlisted_redirect(tmp_path) -> None:
    class Response:
        content = b"%PDF-1.7 official\n" + (b"0" * 256)
        headers = {"content-type": "application/pdf"}
        url = httpx.URL("https://example.com/redirected.pdf")

        def raise_for_status(self):
            return None

    class Client:
        def get(self, _url):
            return Response()

    with pytest.raises(ValueError, match="SOURCE_HOST_NOT_ALLOWED"):
        _validated_direct_content(
            Client(),
            url="https://datafiles.chinhphu.vn/cpp/files/test.pdf",
            cache_dir=tmp_path,
        )


def test_direct_official_gateway_download_uses_encoded_filename_for_validation(
    tmp_path,
) -> None:
    requested_url = (
        "https://vbpl-bientap-gateway.moj.gov.vn/api/qtdc/public/doc/minio/"
        "buckets/vbpl/167031/VanBanGoc_Th%C3%B4ng-tu-08.pdf/download"
    )

    class Response:
        content = b"%PDF-1.7 official\n" + (b"0" * 256)
        headers = {"content-type": "application/octet-stream"}
        url = httpx.URL(requested_url)

        def raise_for_status(self):
            return None

    class Client:
        def get(self, url):
            assert url == requested_url
            return Response()

    content, final_url, file_format = _validated_direct_content(
        Client(),
        url=requested_url,
        cache_dir=tmp_path,
    )

    assert content.startswith(b"%PDF")
    assert final_url == requested_url
    assert file_format == "pdf"


def test_direct_official_gateway_download_uses_query_filename_for_validation(
    tmp_path,
) -> None:
    requested_url = (
        "https://g7.cdnchinhphu.vn/api/download/stream?"
        "token=official&file_name=2025_811_142-2025-N%C4%90-CP.pdf"
    )

    class Response:
        content = b"%PDF-1.7 official\n" + (b"0" * 256)
        headers = {"content-type": "application/pdf"}
        url = httpx.URL(requested_url)

        def raise_for_status(self):
            return None

    class Client:
        def get(self, url):
            assert url == requested_url
            return Response()

    _content, _final_url, file_format = _validated_direct_content(
        Client(),
        url=requested_url,
        cache_dir=tmp_path,
    )

    assert file_format == "pdf"


def test_group_artifact_uses_allowlisted_direct_official_pdf(
    tmp_path, monkeypatch
) -> None:
    official_url = "https://datafiles.chinhphu.vn/cpp/files/96.signed.pdf"
    expected = {
        "file_name": "form.pdf",
        "local_path": "form.pdf",
        "download_url": official_url,
        "sha256": "a" * 64,
        "size_bytes": 512,
        "file_format": "pdf",
    }

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._validated_direct_content",
        lambda *_args, **_kwargs: (b"%PDF-1.7\n" + b"0" * 256, official_url, "pdf"),
    )

    def fake_extract(
        _content,
        *,
        form_code,
        appendix_identifier,
        download_url,
        download_dir,
        allow_ocr=True,
        **_kwargs,
    ):
        assert form_code == "03"
        assert appendix_identifier == "IV"
        assert download_url == official_url
        assert download_dir == tmp_path
        assert allow_ocr is True
        return expected

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources._extract_from_pdf_package",
        fake_extract,
    )

    artifact, reason = resolve_group_artifact(
        {
            "form_code": "03",
            "appendix_identifier": "IV",
            "official_download_urls": [official_url],
        },
        {"document": {"id": "168128"}, "files": []},
        timeout=1,
        download_dir=tmp_path,
    )

    assert artifact is not None
    for key, value in expected.items():
        assert artifact[key] == value
    assert artifact["source_page_url"] == official_url
    assert artifact["source_retrieval_url"] == official_url
    assert artifact["source_retrieval_method"] == "GET"
    assert artifact["source_package_size_bytes"] == 265
    assert artifact["provenance_kind"] == "official_direct_attachment"
    assert reason == "EXTRACTED_FROM_DIRECT_OFFICIAL_PACKAGE"


def test_official_document_search_uses_full_exact_instrument(
    tmp_path, monkeypatch
) -> None:
    class Response:
        def __init__(self, content: bytes):
            self.content = content
            self.status_code = 200
            self.headers = {"content-type": "text/x-component"}

        def raise_for_status(self):
            return None

    class Client:
        payloads: list[list[dict]] = []

        def __init__(self, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def post(self, _url, *, headers, content):
            del headers
            payload = json.loads(content)
            self.payloads.append(payload)
            if len(self.payloads) == 1:
                return Response(
                    b'1:{"total":1,"items":[{"id":"doc-1",'
                    b'"docNum":"02/2024/TT-BYT","title":"Official"}]}\n'
                )
            return Response(b"1:[]\n")

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.httpx.Client", Client
    )

    result = fetch_official_document(
        "02/2024/TT-BYT",
        timeout=1,
        cache_dir=tmp_path,
        refresh=True,
    )

    assert result["status"] == "found"
    assert Client.payloads[0][0]["keyword"] == "02/2024/TT-BYT"


def test_official_document_search_checks_bounded_following_pages(
    tmp_path, monkeypatch
) -> None:
    class Response:
        def __init__(self, content: bytes):
            self.content = content
            self.status_code = 200
            self.headers = {"content-type": "text/x-component"}

        def raise_for_status(self):
            return None

    class Client:
        search_pages: list[int] = []

        def __init__(self, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def post(self, _url, *, headers, content):
            del headers
            payload = json.loads(content)
            if isinstance(payload[0], dict):
                page_number = payload[0]["pageNumber"]
                self.search_pages.append(page_number)
                if page_number == 0:
                    return Response(
                        b'1:{"total":101,"items":[{"id":"other",'
                        b'"docNum":"130/2021/ND-CP","title":"Other"}]}\n'
                    )
                return Response(
                    b'1:{"total":101,"items":[{"id":"doc-131",'
                    b'"docNum":"131/2021/ND-CP","title":"Official"}]}\n'
                )
            return Response(b"1:[]\n")

    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.httpx.Client", Client
    )

    result = fetch_official_document(
        "131/2021/ND-CP",
        timeout=1,
        cache_dir=tmp_path,
        refresh=True,
    )

    assert result["status"] == "found"
    assert result["document"]["id"] == "doc-131"
    assert Client.search_pages == [0, 1]


def test_curated_document_bypasses_unindexed_vbpl_search(tmp_path, monkeypatch) -> None:
    curated = {
        "instrument": "154/2024/NĐ-CP",
        "status": "found",
        "reason_code": "EXACT_CURATED_GOVERNMENT_DOCUMENT_FOUND",
        "checked_url": "https://congbao.chinhphu.vn/van-ban/154.htm",
        "document": {
            "id": "vanban-chinhphu-211821",
            "docNum": "154/2024/NĐ-CP",
            "detailUrl": "https://vanban.chinhphu.vn/?docid=211821",
            "effFrom": "2025-01-10",
            "effStatus": "Còn hiệu lực",
        },
        "files": [],
        "direct_download_urls": [
            "https://congbao.cdnchinhphu.vn/154-2024.pdf"
        ],
    }
    monkeypatch.setattr(
        "scripts.resolve_three_tier_form_sources.load_curated_government_document",
        lambda instrument, *, legal_as_of: curated
        if (instrument, legal_as_of) == ("154/2024/NĐ-CP", "2026-07-30")
        else None,
    )

    result = fetch_official_document(
        "154/2024/NĐ-CP",
        timeout=1,
        cache_dir=tmp_path,
        refresh=True,
        legal_as_of="2026-07-30",
    )

    assert result == curated


def test_official_adapter_order_is_allowlisted_and_deterministic() -> None:
    urls = ordered_fallback_urls(
        [
            "https://example.com/not-official",
            "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/test",
            "https://vbpl.vn/van-ban/02/2024/TT-BYT",
            "https://vbpl.vn/van-ban/02/2024/TT-BYT",
        ],
        instrument="02/2024/TT-BYT",
    )
    assert urls == [
        "https://vbpl.vn/van-ban/02/2024/TT-BYT",
        "https://dichvucong.gov.vn/thu-tuc-hanh-chinh/test",
    ]


def test_official_source_transient_failures_have_resume_safe_reasons() -> None:
    request = httpx.Request("GET", "https://vbpl.vn/van-ban/test")
    assert classify_official_source_failure(
        httpx.ConnectTimeout("timeout", request=request)
    ) == "OFFICIAL_SOURCE_CONNECT_TIMEOUT"
    assert classify_official_source_failure(
        httpx.ReadTimeout("timeout", request=request)
    ) == "OFFICIAL_SOURCE_READ_TIMEOUT"
