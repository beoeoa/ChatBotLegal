from __future__ import annotations

from datetime import date
from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException, Request
from openpyxl import load_workbook

from api.routers import legal_search
from scripts.legal_search_server import (
    LegalRetriever,
    _document_export_workbook,
)


def _request(role: str = "admin") -> Request:
    request = Request(
        {"type": "http", "method": "GET", "path": "/api/legal/documents", "headers": []}
    )
    request.state.user_role = role
    request.state.user_id = f"user:{role}"
    return request


def _bare_retriever() -> LegalRetriever:
    retriever = object.__new__(LegalRetriever)
    retriever._serving_scope = None
    retriever._serving_allowed_document_ids = None
    return retriever


def test_document_browse_filter_is_inclusive_and_manifest_bound(monkeypatch):
    monkeypatch.delenv("LEGAL_SERVING_MANIFEST_REQUIRED", raising=False)
    retriever = _bare_retriever()
    retriever._serving_allowed_document_ids = {9, 3}

    where_sql, params, tier, serving_ids = retriever._document_browse_context(
        query="hộ tịch",
        domain="ho_tich_chung_thuc",
        tier="core",
        as_of=date(2026, 8, 16),
        audience="admin",
        issued_from=date(2026, 1, 1),
        issued_to=date(2026, 8, 16),
    )

    assert "d.issued_date >= :issued_from" in where_sql
    assert "d.issued_date <= :issued_to" in where_sql
    assert "d.id = ANY(:serving_document_ids)" in where_sql
    assert params["serving_document_ids"] == [3, 9]
    assert params["query"] == "%hộ tịch%"
    assert tier == "core"
    assert serving_ids == {3, 9}


def test_document_browse_filter_rejects_reversed_ranges(monkeypatch):
    monkeypatch.delenv("LEGAL_SERVING_MANIFEST_REQUIRED", raising=False)
    retriever = _bare_retriever()

    with pytest.raises(ValueError, match="invalid_effective_date_range"):
        retriever._document_browse_context(
            query="",
            domain=None,
            tier="all",
            as_of=date(2026, 8, 16),
            audience="admin",
            effective_from=date(2026, 8, 17),
            effective_to=date(2026, 8, 16),
        )


def test_document_export_workbook_has_dates_filters_and_formula_protection():
    payload = _document_export_workbook(
        [
            {
                "doc_id": 42,
                "law_number": "=2+2",
                "document_title": "+DDE",
                "document_type": "Nghị định",
                "issuing_agency": "@evil",
                "domain_name": "Hộ tịch",
                "retrieval_tier": "core",
                "issued_date": date(2026, 1, 2),
                "effective_date": "2026-02-03",
                "expired_date": None,
                "effective_status": "active",
                "article_count": 12,
                "source_url": "https://vbpl.vn/example",
            }
        ]
    )
    workbook = load_workbook(BytesIO(payload), data_only=False)
    sheet = workbook.active

    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions
    assert sheet.max_row == 2
    assert sheet["C2"].value == "'=2+2"
    assert sheet["D2"].value == "'+DDE"
    assert sheet["F2"].value == "'@evil"
    assert sheet["I2"].value.date() == date(2026, 1, 2)
    assert sheet["J2"].value.date() == date(2026, 2, 3)
    assert sheet["J2"].number_format == "dd/mm/yyyy"


@pytest.mark.asyncio
async def test_document_list_proxy_forwards_date_filters_and_server_audience(monkeypatch):
    upstream = AsyncMock(
        return_value={
            "items": [],
            "total": 0,
            "limit": 20,
            "offset": 20,
            "as_of": "2026-08-16",
            "tier": "all",
        }
    )
    monkeypatch.setattr(legal_search, "_request", upstream)

    await legal_search.list_legal_documents(
        _request("admin"),
        q="cư trú",
        domain="cu_tru_an_ninh",
        tier="all",
        as_of="2026-08-16",
        issued_from=None,
        issued_to=None,
        effective_from=date(2026, 1, 1),
        effective_to=date(2026, 8, 16),
        expired_from=None,
        expired_to=None,
        limit=20,
        offset=20,
        sort_by="effective_date",
        sort_order="desc",
    )

    params = upstream.await_args.kwargs["params"]
    assert params["effective_from"] == "2026-01-01"
    assert params["effective_to"] == "2026-08-16"
    assert params["audience"] == "admin"
    assert params["limit"] == 20
    assert params["offset"] == 20


@pytest.mark.asyncio
async def test_document_list_proxy_rejects_reversed_range_without_upstream_call(monkeypatch):
    upstream = AsyncMock()
    monkeypatch.setattr(legal_search, "_request", upstream)

    with pytest.raises(HTTPException) as caught:
        await legal_search.list_legal_documents(
            _request(),
            effective_from=date(2026, 8, 17),
            effective_to=date(2026, 8, 16),
        )

    assert caught.value.status_code == 422
    upstream.assert_not_awaited()
