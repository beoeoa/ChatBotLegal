"""Read bounded text from an already-approved local form; no network/parser daemon."""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
from functools import lru_cache
from pathlib import Path
from typing import Any

from api.administrative_query_signals import administrative_domain, fold
from api.official_source_adapters import is_allowlisted_official_url


@lru_cache(maxsize=8)
def _extract(data: bytes, extension: str) -> str:
    if extension == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages[:5])[:10000]
    if extension == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs]
        paragraphs.extend(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        return "\n".join(paragraphs)[:10000]
    return ""


def approved_form_text(
    raw: dict[str, Any],
    root: Path,
    *,
    sidecar_root: Path | None = None,
) -> str:
    """Caller must have passed role/date/binding catalog gates for this form."""
    local, checksum = str(raw.get("local_path") or ""), str(raw.get("sha256") or "").lower()
    if not local or not checksum:
        return ""
    path = (root / local).resolve()
    if not path.is_relative_to(root.resolve()) or not path.is_file() or path.stat().st_size > 5 * 1024 * 1024:
        return ""
    data = path.read_bytes()
    if hashlib.sha256(data).hexdigest() != checksum:
        return ""
    if str(raw.get("file_format") or "").lower().lstrip(".") == "doc":
        # Read-only offline extraction, bound to these exact approved bytes.
        # No desktop automation or optional native process in the API path.
        # Release assets live below ``release-data`` while checksum-bound text
        # sidecars are shared repository artifacts below ``data``. Keep both
        # roots explicit so a valid .doc release can be grounded without
        # copying or trusting an unverified conversion.
        derived_root = (sidecar_root or root).resolve()
        sidecar = derived_root / "data" / "derived" / "form_text" / f"{checksum}.json"
        try:
            if not sidecar.is_file() or sidecar.stat().st_size > 1024 * 1024:
                return ""
            parsed = json.loads(sidecar.read_text(encoding="utf-8"))
            text = parsed.get("text", "")
            if (parsed.get("source_sha256") == checksum and isinstance(text, str)
                    and hashlib.sha256(text.encode("utf-8")).hexdigest() == parsed.get("text_sha256")):
                return text[:10000].strip()
        except (OSError, ValueError, TypeError):
            pass
        return ""
    try:
        return _extract(data, str(raw.get("file_format") or "").lower().lstrip(".")).strip()
    except Exception:
        # Optional parser/invalid text must not take down the legal answer.
        return ""


def build_form_text_evidence(plan, catalog, resolved, *, role, as_of):
    rows = []
    items = {item.item_id: item for item in plan.items}
    for form in (resolved or {}).get("recommended_forms") or []:
        raw = next((f for f in catalog.forms if f.get("form_id") == form.get("form_id")), None)
        if raw is None or not catalog.gate_form(raw, procedure_id=form["procedure_id"], role=role, as_of=as_of).eligible:
            continue
        requested = [i for i in form.get("issue_ids", []) if i in items and
                     (items[i].action == "form_lookup" or set(items[i].facets) & {"form", "consent", "documents"})]
        if not requested:
            continue
        text = approved_form_text(raw, catalog.project_root)
        if not text:
            continue
        for issue_id in requested:
            rows.append({
                "source_id": f"form-text:{form['form_id']}:{issue_id}", "issue_id": issue_id,
                "document_title": form.get("display_name"), "source_url": form.get("source_url"),
                "form_id": form["form_id"], "procedure_id": form["procedure_id"],
                "supported_facets": ["form", "consent"], "effective_from": form.get("effective_from"),
                "effective_to": form.get("effective_to"), "source_checksum": raw["sha256"],
                "evidence_kind": "approved_form_text", "content": "Nội dung trích từ biểu mẫu đã duyệt (không chứng minh yêu cầu phải đi cùng):\n" + text,
            })
    return rows


def build_active_release_form_text_evidence(plan, resolved, *, role, as_of):
    """Read form text only from assets in the current active release.

    This mirrors ``build_form_text_evidence`` but binds bytes to the live
    release pointer.  It prevents a withdrawn/replaced form from surviving in
    chatbot context through the legacy JSON catalog cache.
    """

    del role, as_of  # Audience/date eligibility was already enforced by resolver.
    from api.form_governance_service import get_form_governance_service

    release = get_form_governance_service().repository.active_release()
    if not release:
        return []
    manifest = release.get("manifest") or release
    assets = {
        str(item.get("form_id") or ""): item
        for item in manifest.get("assets") or []
        if isinstance(item, dict) and item.get("coverage_status") == "released"
    }
    runtime_root = Path(
        str(
            os.getenv("FORM_RELEASE_ASSET_ROOT")
            or Path(__file__).resolve().parents[1] / "release-data"
        )
    ).resolve()
    items = {item.item_id: item for item in plan.items}
    rows = []
    for form in (resolved or {}).get("recommended_forms") or []:
        form_id = str(form.get("form_id") or "")
        raw = assets.get(form_id)
        if not raw or raw.get("asset_kind") != "file":
            continue
        runtime_path = str(raw.get("runtime_path") or "").strip()
        checksum = str(raw.get("source_checksum") or "").casefold()
        if (
            not runtime_path
            or not checksum
            or raw.get("download_url")
            != f"/api/procedures/forms-catalog/assets/{form_id}/download"
        ):
            continue
        requested = [
            issue_id
            for issue_id in form.get("issue_ids", [])
            if issue_id in items
            and (
                items[issue_id].action == "form_lookup"
                or set(items[issue_id].facets) & {"form", "consent", "documents"}
            )
        ]
        if not requested:
            continue
        text = approved_form_text(
            {
                "local_path": runtime_path,
                "sha256": checksum,
                "file_format": raw.get("file_format") or raw.get("file_type"),
            },
            runtime_root,
            sidecar_root=Path(__file__).resolve().parents[1],
        )
        if not text:
            continue
        for issue_id in requested:
            rows.append(
                {
                    "source_id": f"form-release-text:{form_id}:{issue_id}",
                    "issue_id": issue_id,
                    "document_title": form.get("display_name") or form.get("name"),
                    "source_url": form.get("source_url"),
                    "form_id": form_id,
                    "procedure_id": form.get("procedure_id"),
                    "supported_facets": ["form", "consent"],
                    "effective_from": form.get("effective_from"),
                    "effective_to": form.get("effective_to"),
                    "source_checksum": checksum,
                    "evidence_kind": "active_release_form_text",
                    "content": (
                        "Nội dung trích từ biểu mẫu trong bản phát hành đang hoạt động "
                        "(không tự chứng minh yêu cầu phải đi cùng):\n" + text
                    ),
                }
            )
    return rows


def build_active_release_form_comparison_evidence(
    question: str,
    *,
    issue_id: str,
    role: str,
    as_of,
) -> list[dict[str, Any]]:
    """Project exact, checksum-bound metadata for an explicit form comparison."""

    query = fold(question)
    if not any(marker in query for marker in ("phan biet", "khac nhau", "de tranh nham", "de khong nham")):
        return []
    codes: list[str] = []
    for match in re.finditer(r"\bct\s*0?(\d{1,3})\b", query):
        code = f"CT{int(match.group(1)):02d}"
        if code not in codes:
            codes.append(code)
    for match in re.finditer(r"\bmau\s+(?:so\s+)?(\d{1,3})\b", query):
        code = str(int(match.group(1))).zfill(2)
        if code not in codes:
            codes.append(code)
    if len(codes) < 2:
        return []

    from api.form_governance_service import get_form_governance_service
    from api.form_router_v3 import _audience_is_allowed, _effective

    release = get_form_governance_service().repository.active_release()
    if not release:
        return []
    manifest = release.get("manifest") or release
    assets = {
        str(item.get("form_id") or ""): item
        for item in manifest.get("assets") or []
        if isinstance(item, dict)
    }
    procedures = {
        str(item.get("procedure_id") or ""): item
        for item in manifest.get("procedures") or []
        if isinstance(item, dict)
    }
    bindings = [
        item for item in manifest.get("bindings") or []
        if isinstance(item, dict)
    ]
    wanted_domain = administrative_domain(question)
    codes_by_procedure: dict[str, set[str]] = {}
    for binding in bindings:
        asset = assets.get(str(binding.get("form_id") or "")) or {}
        code = str(asset.get("form_code") or "").replace(" ", "").upper()
        if code in codes:
            codes_by_procedure.setdefault(
                str(binding.get("procedure_id") or ""), set()
            ).add(code)

    query_tokens = set(query.split())
    rows: list[dict[str, Any]] = []
    used_assets: set[str] = set()
    for code in codes:
        candidates: list[tuple[int, str, dict[str, Any], dict[str, Any], dict[str, Any]]] = []
        for binding in bindings:
            form_id = str(binding.get("form_id") or "")
            asset = assets.get(form_id)
            procedure = procedures.get(str(binding.get("procedure_id") or ""))
            if not asset or not procedure:
                continue
            actual_code = str(asset.get("form_code") or "").replace(" ", "").upper()
            if actual_code != code:
                continue
            if (
                str(binding.get("coverage_status") or "") != "released"
                or str(asset.get("coverage_status") or "") != "released"
                or str(procedure.get("coverage_status") or "") != "released"
                or not _audience_is_allowed(binding.get("audience"), role)
                or not _audience_is_allowed(asset.get("audiences"), role)
                or not _effective(binding, as_of)
                or not _effective(asset, as_of)
                or not _effective(procedure, as_of)
                or not is_allowlisted_official_url(str(asset.get("source_url") or ""))
                or not re.fullmatch(r"[0-9a-f]{64}", str(asset.get("source_checksum") or "").casefold())
            ):
                continue
            procedure_domain = str(procedure.get("domain") or "")
            if wanted_domain and procedure_domain != wanted_domain:
                continue
            label = fold(
                f"{asset.get('canonical_name') or ''} "
                f"{procedure.get('name') or procedure.get('canonical_name') or ''}"
            )
            overlap = len(query_tokens & set(label.split()))
            shared = len(codes_by_procedure.get(str(binding.get("procedure_id") or ""), set()))
            score = shared * 100 + overlap
            candidates.append((score, form_id, asset, binding, procedure))
        if not candidates:
            continue
        _, form_id, asset, binding, procedure = max(
            candidates, key=lambda item: (item[0], item[1])
        )
        if form_id in used_assets:
            continue
        used_assets.add(form_id)
        legal_basis = binding.get("legal_basis") or asset.get("issuing_instrument")
        if isinstance(legal_basis, list):
            legal_basis_text = "; ".join(str(value) for value in legal_basis if value)
        else:
            legal_basis_text = str(legal_basis or "")
        verified_as_of = str(manifest.get("legal_as_of") or as_of.isoformat())
        content = (
            "Metadata biểu mẫu trong bản phát hành đang hoạt động:\n"
            f"- Mã biểu mẫu: {asset.get('form_code')}\n"
            f"- Tên chính thức: {asset.get('canonical_name')}\n"
            f"- Thủ tục được gắn: {procedure.get('name') or procedure.get('canonical_name')}\n"
            f"- Căn cứ ban hành: {legal_basis_text or 'chưa có số hiệu riêng trong metadata'}\n"
            f"- Đủ điều kiện phát hành tại ngày đối chiếu: {verified_as_of}."
        )
        rows.append({
            "source_id": f"form-comparison:{form_id}:{issue_id}",
            "issue_id": issue_id,
            "document_title": str(asset.get("canonical_name") or ""),
            "source_url": str(asset.get("source_url") or ""),
            "form_id": form_id,
            "procedure_id": str(procedure.get("procedure_id") or ""),
            "supported_facets": ["form", "documents"],
            "effective_from": asset.get("effective_from"),
            "effective_to": asset.get("effective_to"),
            "verified_as_of": verified_as_of,
            "source_checksum": str(asset.get("source_checksum") or ""),
            "evidence_kind": "active_release_form_comparison",
            "content": content,
        })
    return rows if len(rows) == len(codes) else []
